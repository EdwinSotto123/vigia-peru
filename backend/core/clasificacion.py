"""Clasificación tipo × etapa de un proceso SEACE y agentes que le aplican.

Entrada: un release OCDS (recortado como el `convocatorias.ocds_payload` que guarda
`backend/scrapers/oece_ocds`, o el `compiledRelease` completo de `/record/<ocid>`).
Salida: `Clasificacion` (tipo, etapa, modalidad, procesable, agentes, validaciones
pendientes, RUC del proveedor). La matriz completa está en
`docs/design/MATRIZ_TIPO_ETAPA.md`.

Hechos medidos (2026-09-15, 18 413 releases jun→sep 2026) que definen las reglas:
  · `tender.status` es SIEMPRE null → la etapa NO puede salir de ahí.
  · `tag` acumula la composición del release: planning/tender/award/contract/implementation
    (y `compiled` en los records completos).
  · `tender.items[].statusDetails` es la señal de etapa más fina del SEACE: CONVOCADO,
    CONSENTIDO, ADJUDICADO, CONTRATADO, DESIERTO, NULO, CANCELADO, APELADO, SUSPENDIDO,
    RETROTRAIDO_POR_RESOLUCION, PENDIENTE_DE_REGISTRO_DE_EFECTO, NO_SUSCRIPCION_CONTRATO…
    (desierta/nula/cancelada solo se ven ahí en los releases recortados).
  · El release recortado NO trae `awards`/`contracts`/`parties` → el RUC del proveedor
    solo se conoce con el record completo; el orquestador lo obtiene al analizar.
  · Ningún release real trae tag solo-`planning`: SEACE publica planning+tender juntos.

Uso como script (recorre `convocatorias` y escribe las columnas de la migración 13):
  python -m backend.core.clasificacion --reclasificar --dry-run
  python -m backend.core.clasificacion --reclasificar [--lote 1000] [--solo-sin-clasificar]
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import asdict, dataclass, field
from pathlib import Path

# ── Vocabulario ─────────────────────────────────────────────────────────────
TIPOS = ("bienes", "servicios", "consultoria", "obras", "convenio", "directa", "otro")
ETAPAS = ("planificacion", "convocada", "adjudicada", "contratada", "en_ejecucion", "finalizada",
          "desierta", "cancelada", "nula", "desconocida")
AGENTES = ("compliance", "document_parser", "document_legal_analyst", "market", "web_research",
           "news_research", "entity_personnel", "person_network", "compliance_extended", "report_writer")

ETAPAS_POST_ADJUDICACION = ("adjudicada", "contratada", "en_ejecucion", "finalizada")
ETAPAS_CON_CONTRATO = ("contratada", "en_ejecucion", "finalizada")
ETAPAS_NEGATIVAS = ("desierta", "cancelada", "nula")

# Motivos de `procesable = false`.
MOTIVO_TIPO_NO_SOPORTADO = "tipo_no_soportado"
MOTIVO_ETAPA_DESCONOCIDA = "etapa_desconocida"
MOTIVO_ETAPA_PLANIFICACION = "etapa_planificacion"   # sin postores ni bases: nada que analizar aún

# Validaciones pendientes (códigos estables; el frontend y el dictamen los describen).
VALIDACIONES = {
    "infobras_avance": "Avance físico/financiero de la obra (INFOBRAS) no disponible: no se verificó la ejecución.",
    "market_sin_items_fisicos": "Sin ítems físicos con cantidad y unidad: no se comparó precio de mercado (solo bienes).",
    "sin_documentos_descargables": "Sin bases, buena pro ni contrato publicados: no se analizaron documentos.",
    "proveedor_sin_ruc": "El record no identifica al proveedor (RUC): no se investigó a la empresa ni su red.",
    "entidad_sin_ruc": "Sin RUC de la entidad: no se investigó a sus funcionarios.",
}

# Tipos de documento que alimentan al document_parser (los demás — clarifications,
# evaluationReports — no aportan el requerimiento ni el contrato).
_DOCS_CLAVE = {"biddingDocuments", "awardNotice", "contractSigned"}

# statusDetails de ítems → familia de etapa.
_ITEMS_CONTRATADO = {"CONTRATADO"}
_ITEMS_ADJUDICADO = {"ADJUDICADO", "CONSENTIDO", "APELADO"}
_ITEMS_NULO = {"NULO"}
_ITEMS_CANCELADO = {"CANCELADO", "NO_SUSCRIPCION_CONTRATO POR_DECISION_ENTIDAD", "NO_SUSCRIPCION_CONTRATO_POR_DECISION_ENTIDAD"}
_ITEMS_DESIERTO = {"DESIERTO"}
_ITEMS_NEGATIVOS = _ITEMS_NULO | _ITEMS_CANCELADO | _ITEMS_DESIERTO


@dataclass
class Clasificacion:
    tipo: str
    etapa: str
    modalidad: str | None
    procesable: bool
    motivo_no_procesable: str | None
    agentes: list[str] = field(default_factory=list)
    validaciones_pendientes: list[str] = field(default_factory=list)
    proveedor_ruc: str | None = None

    def as_dict(self) -> dict:
        return asdict(self)


# ── Matriz tipo × etapa → agentes ───────────────────────────────────────────
def _ordenar(agentes) -> list[str]:
    """Orden canónico del pipeline (el mismo de deterministic.py / events.FASES)."""
    s = set(agentes)
    return [a for a in AGENTES if a in s]


def _construir_matriz() -> dict[tuple[str, str], list[str]]:
    # `market` corre en TODOS los tipos: la estrategia la decide el perfil del servicio
    # (`agents/_shared/profiles.py` → `market_estrategia`: goods_retail para bienes,
    # historico_seace para servicios/consultoría, presupuesto_obra para obras, cotizaciones
    # para convenio/directa). La matriz ya no lo omite a ciegas para no-bienes.
    base = ["compliance", "document_parser", "document_legal_analyst", "market", "entity_personnel", "report_writer"]
    investigacion = ["web_research", "news_research", "person_network", "compliance_extended"]
    base_convenio = ["compliance", "document_parser", "document_legal_analyst", "market", "report_writer"]
    # En las etapas negativas el dictamen es breve (causal + contexto). Se incluye report_writer
    # también para convenio/directa: una corrida sin dictamen deja una alerta ilegible.
    negativas = ["compliance", "report_writer"]
    m: dict[tuple[str, str], list[str]] = {}
    for tipo in TIPOS:
        if tipo == "otro":
            continue
        m[(tipo, "planificacion")] = []
        for etapa in ETAPAS_NEGATIVAS:
            m[(tipo, etapa)] = list(negativas)
        if tipo in ("bienes", "servicios", "consultoria", "obras"):
            m[(tipo, "convocada")] = _ordenar(base)
            for etapa in ETAPAS_POST_ADJUDICACION:
                m[(tipo, etapa)] = _ordenar(base + investigacion)       # los 10
        else:  # convenio · directa
            m[(tipo, "convocada")] = _ordenar(base_convenio)
            m[(tipo, "adjudicada")] = _ordenar(base_convenio + ["person_network", "web_research", "news_research"])
            for etapa in ETAPAS_CON_CONTRATO:
                m[(tipo, etapa)] = list(AGENTES)
    return m


MATRIZ: dict[tuple[str, str], list[str]] = _construir_matriz()


# ── Helpers de lectura del OCDS ─────────────────────────────────────────────
def _norm(s) -> str:
    """minúsculas sin tildes ('Regímen Especial' → 'regimen especial')."""
    s = unicodedata.normalize("NFKD", str(s or ""))
    return "".join(c for c in s if not unicodedata.combining(c)).lower().strip()


def _ocid_corto(ocid: str) -> str:
    """'ocds-dgv273-seacev3-1212353' → '1212353' (misma regla que oece_ocds.short_ocid)."""
    s = str(ocid or "")
    return s[len("ocds-dgv273-seacev3-"):] if s.startswith("ocds-dgv273-seacev3-") else s


def _ruc_de_id(ident) -> str | None:
    """'PE-RUC-20508320664' → '20508320664'; acepta el RUC pelado."""
    m = re.search(r"(\d{11})$", str(ident or ""))
    return m.group(1) if m else None


def _parties(rel: dict) -> list[dict]:
    return [p for p in (rel.get("parties") or []) if isinstance(p, dict)]


def _ruc_de_party(p: dict) -> str | None:
    for ai in p.get("additionalIdentifiers") or []:
        if isinstance(ai, dict) and ai.get("scheme") == "PE-RUC" and _ruc_de_id(ai.get("id")):
            return _ruc_de_id(ai.get("id"))
    ident = p.get("identifier") or {}
    if ident.get("scheme") == "PE-RUC" and _ruc_de_id(ident.get("id")):
        return _ruc_de_id(ident.get("id"))
    return None


def entidad_ruc(rel: dict) -> str | None:
    """RUC de la entidad contratante (parties con rol buyer/procuringEntity)."""
    for p in _parties(rel):
        roles = p.get("roles") or []
        if "buyer" in roles or "procuringEntity" in roles:
            r = _ruc_de_party(p)
            if r:
                return r
    return None


def _awards_activos(rel: dict) -> list[dict]:
    """Awards no fallidos (status null = activo en los records del OECE)."""
    return [a for a in (rel.get("awards") or []) if isinstance(a, dict)
            and _norm(a.get("status") or "active") not in ("unsuccessful", "cancelled")]


def proveedor_ruc(rel: dict) -> str | None:
    """RUC del proveedor adjudicado: awards[].suppliers[].id o party con rol supplier."""
    for a in _awards_activos(rel):
        for s in a.get("suppliers") or []:
            r = _ruc_de_id((s or {}).get("id")) if isinstance(s, dict) else None
            if r:
                return r
    for p in _parties(rel):
        if "supplier" in (p.get("roles") or []):
            r = _ruc_de_party(p)
            if r:
                return r
    return None


def _items(rel: dict) -> list[dict]:
    return [i for i in ((rel.get("tender") or {}).get("items") or []) if isinstance(i, dict)]


def _items_status(rel: dict) -> set[str]:
    out = set()
    for i in _items(rel):
        st = str(i.get("statusDetails") or "").upper().strip()
        if st:
            out.add(st)
    return out


def _implementacion_activa(c: dict) -> bool:
    impl = c.get("implementation")
    if not isinstance(impl, dict):
        return False
    return bool(impl.get("transactions") or impl.get("milestones"))


def documentos(rel: dict) -> list[dict]:
    """Todos los documentos publicados (tender + awards + contracts)."""
    out: list[dict] = []
    for d in (rel.get("tender") or {}).get("documents") or []:
        if isinstance(d, dict):
            out.append(d)
    for grupo in ("awards", "contracts"):
        for x in rel.get(grupo) or []:
            if not isinstance(x, dict):
                continue
            for d in x.get("documents") or []:
                if isinstance(d, dict):
                    out.append(d)
    return out


def _hay_documentos_clave(rel: dict) -> bool:
    return any(d.get("documentType") in _DOCS_CLAVE and d.get("url") for d in documentos(rel))


def _item_fisico(item: dict, categoria: str) -> bool:
    """Ítem comparable en mercado: cantidad > 0, unidad, y bien físico (categoría goods o
    clasificación CUBSO/UNSPSC de producto: segmento < 70)."""
    try:
        q = float(item.get("quantity") or 0)
    except (TypeError, ValueError):
        q = 0.0
    if q <= 0 or not (item.get("unit") or {}).get("name"):
        return False
    if categoria == "goods":
        return True
    ids = [(item.get("classification") or {}).get("id")]
    ids += [(c or {}).get("id") for c in item.get("additionalClassifications") or []]
    for cid in ids:
        cid = str(cid or "")
        if len(cid) >= 2 and cid[:2].isdigit() and int(cid[:2]) < 70:
            return True
    return False


# ── Derivación de tipo y etapa ──────────────────────────────────────────────
def derivar_tipo(rel: dict) -> str:
    """`mainProcurementCategory` + `procurementMethodDetails` → tipo. La modalidad manda
    cuando cambia el régimen (directa, convenio) o la naturaleza (consultoría)."""
    tender = rel.get("tender") or {}
    cat = _norm(tender.get("mainProcurementCategory"))
    mod = _norm(tender.get("procurementMethodDetails"))
    texto = _norm(f"{tender.get('title') or ''} {tender.get('description') or ''}")
    if "contratacion directa" in mod:
        return "directa"
    if mod.startswith("convenio") or "regimen especial" in mod or "contratacion internacional" in mod:
        return "convenio"
    if "consultoria" in mod or "consultoria de obra" in texto:
        return "consultoria"
    return {"goods": "bienes", "services": "servicios", "works": "obras"}.get(cat, "otro")


def derivar_etapa(rel: dict, *, nulo: bool = False) -> str:
    """Etapa a partir de objetos del record (contracts/awards), tags del release y
    `tender.items[].statusDetails` — en ese orden de confianza. `tender.status` casi nunca
    viene, pero se respeta si está."""
    if nulo:
        return "nula"
    tender = rel.get("tender") or {}
    tags = {str(t) for t in (rel.get("tag") or [])}
    contracts = [c for c in (rel.get("contracts") or []) if isinstance(c, dict)]
    awards = [a for a in (rel.get("awards") or []) if isinstance(a, dict)]
    status = _norm(tender.get("status"))
    items = _items_status(rel)

    if any(_norm(c.get("status")) == "terminated" for c in contracts):
        return "finalizada"
    if "implementation" in tags or any(_implementacion_activa(c) for c in contracts):
        return "en_ejecucion"
    if contracts or "contract" in tags or items & _ITEMS_CONTRATADO:
        return "contratada"
    if _awards_activos(rel) or "award" in tags or items & _ITEMS_ADJUDICADO:
        return "adjudicada"
    if status == "cancelled":
        return "cancelada"
    if status == "unsuccessful" or (awards and not _awards_activos(rel)):
        return "desierta"
    if items and items <= _ITEMS_NEGATIVOS:
        # Todos los ítems terminaron sin adjudicar: nula > cancelada > desierta.
        if items & _ITEMS_NULO:
            return "nula"
        if items & _ITEMS_CANCELADO:
            return "cancelada"
        return "desierta"
    if "tender" in tags or tender.get("items") or tender.get("tenderPeriod") or tender.get("datePublished"):
        return "convocada"
    if "planning" in tags or rel.get("planning"):
        return "planificacion"
    return "desconocida"


# ── Clasificación completa ──────────────────────────────────────────────────
def clasificar(release_o_record: dict, *, entidad_ruc_hint: str | None = None, nulo: bool = False) -> Clasificacion:
    """Release recortado o record completo → Clasificacion.

    `entidad_ruc_hint`: RUC ya conocido de la entidad (columna `convocatorias.entidad_ruc`),
    porque el release recortado no guarda `parties`. `nulo`: marcado en `dataset/nulos/`.
    Nunca inventa: si a un agente le faltan datos, se omite y queda en `validaciones_pendientes`.
    """
    rel = release_o_record if isinstance(release_o_record, dict) else {}
    tender = rel.get("tender") or {}
    tipo = derivar_tipo(rel)
    etapa = derivar_etapa(rel, nulo=nulo)
    modalidad = (str(tender.get("procurementMethodDetails")).strip() or None) if tender.get("procurementMethodDetails") else None
    ruc_prov = proveedor_ruc(rel)
    ruc_ent = entidad_ruc(rel) or entidad_ruc_hint

    if tipo == "otro":
        return Clasificacion(tipo, etapa, modalidad, False, MOTIVO_TIPO_NO_SOPORTADO, [], [], ruc_prov)
    if etapa == "desconocida":
        return Clasificacion(tipo, etapa, modalidad, False, MOTIVO_ETAPA_DESCONOCIDA, [], [], ruc_prov)
    if etapa == "planificacion":
        return Clasificacion(tipo, etapa, modalidad, False, MOTIVO_ETAPA_PLANIFICACION, [], [], ruc_prov)

    agentes = list(MATRIZ.get((tipo, etapa), []))
    pendientes: list[str] = []

    def _omitir(agente: str, codigo: str) -> None:
        if agente in agentes:
            agentes.remove(agente)
            if codigo not in pendientes:
                pendientes.append(codigo)

    # Datos requeridos por agente (ver §1 del plan / MATRIZ_TIPO_ETAPA.md).
    # `market_sin_items_fisicos` solo aplica a BIENES (precio unitario × cantidad en retail);
    # en servicios/obras/convenio/directa el mercado se compara por histórico SEACE,
    # presupuesto del expediente o cotizaciones del sustento (no necesita ítems físicos).
    if "market" in agentes and tipo == "bienes":
        cat = _norm(tender.get("mainProcurementCategory"))
        if not any(_item_fisico(i, cat) for i in _items(rel)):
            _omitir("market", "market_sin_items_fisicos")
    if "document_parser" in agentes and not _hay_documentos_clave(rel):
        _omitir("document_parser", "sin_documentos_descargables")
        _omitir("document_legal_analyst", "sin_documentos_descargables")   # analiza lo que el parser extrajo
    # El RUC del proveedor solo es verificable con el record completo (`awards` presente).
    # Con un release recortado en etapa post-adjudicación, el orquestador lo resolverá al
    # traer el record: no se omite a ciegas.
    if "awards" in rel and not ruc_prov:
        for a in ("web_research", "news_research", "person_network", "compliance_extended"):
            _omitir(a, "proveedor_sin_ruc")
    if "entity_personnel" in agentes and not ruc_ent:
        _omitir("entity_personnel", "entidad_sin_ruc")
    if tipo == "obras" and etapa in ETAPAS_CON_CONTRATO:
        pendientes.append("infobras_avance")   # pendiente parcial: no omite agentes

    return Clasificacion(tipo, etapa, modalidad, True, None, _ordenar(agentes), pendientes, ruc_prov)


def describir_validacion(codigo: str) -> str:
    return VALIDACIONES.get(codigo, codigo)


# ── Nulos locales (dataset/nulos/*.xlsx, columna `codigoconvocatoria` = ocid corto) ──
def cargar_nulos(paths: list[Path] | None = None) -> set[str]:
    """OCIDs cortos declarados nulos en los listados CONOSCE locales. Vacío si no hay
    archivos o falta openpyxl (dependencia opcional)."""
    from backend.scrapers._core.storage import REPO_ROOT  # import tardío: solo el script lo usa

    paths = paths if paths is not None else sorted((REPO_ROOT / "dataset" / "nulos").glob("*.xlsx"))
    out: set[str] = set()
    if not paths:
        return out
    try:
        import openpyxl
    except ImportError:
        return out
    for p in paths:
        wb = openpyxl.load_workbook(p, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(rows, [])]
        if "codigoconvocatoria" not in header:
            continue
        col = header.index("codigoconvocatoria")
        for r in rows:
            v = r[col] if col < len(r) else None
            if v not in (None, ""):
                out.add(str(v).strip().split(".")[0])
    return out


# ── Script: reclasificar convocatorias en la DB ──────────────────────────────
_UPDATE_SQL = """
UPDATE convocatorias c SET
  tipo_contratacion = v.tipo, etapa = v.etapa, modalidad = v.modalidad, procesable = v.procesable,
  motivo_no_procesable = v.motivo, agentes_aplicables = v.agentes, validaciones_pendientes = v.validaciones,
  proveedor_ruc = COALESCE(v.proveedor_ruc, c.proveedor_ruc), clasificado_at = now()
FROM (VALUES %s) AS v(ocid, tipo, etapa, modalidad, procesable, motivo, agentes, validaciones, proveedor_ruc)
WHERE c.ocid = v.ocid
"""
_UPDATE_TEMPLATE = "(%s, %s, %s, %s, %s::boolean, %s, %s::text[], %s::text[], %s::char(11))"


def _fila(ocid: str, c: Clasificacion) -> tuple:
    return (ocid, c.tipo, c.etapa, c.modalidad, c.procesable, c.motivo_no_procesable,
            list(c.agentes), list(c.validaciones_pendientes), c.proveedor_ruc)


def reclasificar(dsn: str, *, lote: int = 1000, dry_run: bool = False, solo_sin_clasificar: bool = False,
                 nulos: set[str] | None = None, log=print) -> dict:
    """Recorre `convocatorias` con cursor de servidor y actualiza en lotes.
    Devuelve la distribución {(tipo, etapa, procesable): n} más los motivos."""
    import collections

    import psycopg2
    from psycopg2.extras import execute_values

    nulos = nulos or set()
    dist: collections.Counter = collections.Counter()
    motivos: collections.Counter = collections.Counter()
    validaciones: collections.Counter = collections.Counter()
    where = "WHERE clasificado_at IS NULL" if solo_sin_clasificar else ""
    conn = psycopg2.connect(dsn)
    total = 0
    try:
        # withhold: el cursor de servidor sobrevive a los commit de cada lote.
        with conn.cursor(name="reclasificar", withhold=True) as cur:
            cur.itersize = lote
            cur.execute(f"SELECT ocid, entidad_ruc, ocds_payload FROM convocatorias {where} ORDER BY ocid")
            buf: list[tuple] = []
            while True:
                rows = cur.fetchmany(lote)
                if not rows:
                    break
                for ocid, ruc_ent, payload in rows:
                    c = clasificar(payload or {}, entidad_ruc_hint=ruc_ent, nulo=(_ocid_corto(ocid) in nulos))
                    dist[(c.tipo, c.etapa, c.procesable)] += 1
                    if c.motivo_no_procesable:
                        motivos[c.motivo_no_procesable] += 1
                    for v in c.validaciones_pendientes:
                        validaciones[v] += 1
                    buf.append(_fila(ocid, c))
                total += len(rows)
                if not dry_run:
                    with conn.cursor() as w:
                        execute_values(w, _UPDATE_SQL, buf, template=_UPDATE_TEMPLATE, page_size=lote)
                    conn.commit()
                buf.clear()
                log(f"   {total:6d} filas {'clasificadas (dry-run)' if dry_run else 'actualizadas'}")
    finally:
        conn.close()
    return {"total": total, "distribucion": dict(dist), "motivos": dict(motivos), "validaciones": dict(validaciones)}


def _imprimir_resumen(res: dict, log=print) -> None:
    dist = res["distribucion"]
    log(f"\n→ {res['total']} convocatorias · tipo × etapa × procesable:")
    for (tipo, etapa, proc), n in sorted(dist.items(), key=lambda kv: (-kv[1], kv[0])):
        log(f"   {n:6d}  {tipo:12s} {etapa:14s} {'procesable' if proc else 'PENDIENTE'}")
    por_tipo: dict[str, int] = {}
    por_etapa: dict[str, int] = {}
    proc_n = 0
    for (tipo, etapa, proc), n in dist.items():
        por_tipo[tipo] = por_tipo.get(tipo, 0) + n
        por_etapa[etapa] = por_etapa.get(etapa, 0) + n
        proc_n += n if proc else 0
    log(f"   por tipo: {dict(sorted(por_tipo.items(), key=lambda kv: -kv[1]))}")
    log(f"   por etapa: {dict(sorted(por_etapa.items(), key=lambda kv: -kv[1]))}")
    log(f"   procesables: {proc_n} · pendientes de procesamiento: {res['total'] - proc_n} · motivos: {res['motivos']}")
    log(f"   validaciones pendientes: {res['validaciones']}")


def main(argv: list[str] | None = None) -> int:
    import argparse

    ap = argparse.ArgumentParser(description="Clasificación tipo × etapa de convocatorias SEACE")
    ap.add_argument("--reclasificar", action="store_true", help="recorre convocatorias y escribe las columnas de la migración 13")
    ap.add_argument("--dry-run", action="store_true", help="clasifica y muestra la distribución sin escribir")
    ap.add_argument("--lote", type=int, default=1000)
    ap.add_argument("--solo-sin-clasificar", action="store_true", help="solo filas con clasificado_at IS NULL")
    ap.add_argument("--sin-nulos", action="store_true", help="no cruzar con dataset/nulos/*.xlsx")
    args = ap.parse_args(argv)
    if not args.reclasificar:
        ap.print_help()
        return 2
    from backend.scrapers._core.pipeline import pg_dsn

    nulos = set() if args.sin_nulos else cargar_nulos()
    print(f"→ reclasificar (lote {args.lote}{', dry-run' if args.dry_run else ''}) · {len(nulos)} ocids nulos locales")
    res = reclasificar(pg_dsn(), lote=args.lote, dry_run=args.dry_run, solo_sin_clasificar=args.solo_sin_clasificar, nulos=nulos)
    _imprimir_resumen(res)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
