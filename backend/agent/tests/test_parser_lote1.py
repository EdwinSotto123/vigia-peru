"""Tests OFFLINE del Workstream R4 (correcciones del lote 1 · T7) sobre tools/documentos.py,
tools/doc_select.py y tools/docai.py. Cada test reproduce un caso de
docs/design/revision_contratos/*.md con fixtures sintéticas (sin red, Gemini, Document AI ni BD).

  · 1225058 / 1225266 / 1225090: la marca y el precio de la OC/contrato NO se funden en el ítem
    de las bases ("las bases exigen la marca X" inventado) → items_contratados + cruce.
  · 1225392 / 1225379 / 1225030: postores con RUC repetido se FUSIONAN (montos, ganador, puntaje).
  · 1225090: RUC mal leído (dígito verificador) no crea un tercer postor; se corrige contra el OCDS.
  · 1225379: lista de invitados (CP) y ganador no invitado.
  · 1225266 / 1225030: PDF idéntico 3× en el ZIP → 1 OCR; mismo sha en el lote → 1 OCR.
  · 1225256 / 1225266: imágenes del DOCX en orden del cuerpo (r:embed), páginas reales.
  · 1225416 / 1225058: folio impreso ≠ índice ⟦p.N⟧ → se corrige la página y se anota el folio.
  · 1225058: bloque `contrato` (ampliación improcedente); 1225090: `procedimiento_seleccion`.
"""
from __future__ import annotations

import io
import os
import sys
import zipfile

import pytest

_AGENT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _AGENT not in sys.path:
    sys.path.insert(0, _AGENT)

os.environ.pop("DOCAI_PROCESSOR_ID", None)

import tools.documentos as D  # noqa: E402
from tools.doc_select import seleccionar_documentos, categorias_de, rank_documento  # noqa: E402

PRIO_BIENES = ("Bases integradas", "Bases", "EETT", "Resumen ejecutivo", "Acta/Cuadro comparativo",
               "Contrato/Orden", "Adendas", "Absolución")


# ── helpers ───────────────────────────────────────────────────────────────────────────
def _pdf(textos: list[str]) -> bytes:
    import fitz
    doc = fitz.open()
    for t in textos:
        page = doc.new_page()
        page.insert_text((72, 72), t, fontsize=11)
    b = doc.tobytes()
    doc.close()
    return b


def _zip(entries: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for k, v in entries.items():
            z.writestr(k, v)
    return buf.getvalue()


def _doc(id_, titulo, tipo, sha, seccion="tender", formato="pdf"):
    return {"id": id_, "url": f"https://x/{id_}", "gs": None, "tipo": tipo, "titulo": titulo,
            "seccion": seccion, "formato": formato, "sha256": sha}


def _tx(paginas: list[str]) -> dict:
    pags = [{"n": i + 1, "texto": t, "chars": len(t)} for i, t in enumerate(paginas)]
    return {"n_paginas": len(pags), "chars": sum(p["chars"] for p in pags), "motor": "docai", "truncado": False,
            "paginas": pags}


def _res(doc: dict, ext: dict, paginas: list[str]) -> dict:
    """Resultado de _procesar_doc tal como lo consume la consolidación: la extracción pasa por
    _post_procesar (evidencia + remapeo por etapa), igual que en producción."""
    tx = _tx(paginas)
    ext = D._post_procesar(dict(ext), tx, doc["sha256"], doc)
    ext.setdefault("_usos", [])
    return {"doc": doc, "sha256": doc["sha256"], "recortes": [], "tiempos": {"total_s": 0.1},
            "cache": {"texto": False, "extraccion": False}, "tx": tx, "ext": ext}


def _lote(monkeypatch, docs, fake: dict, state: dict | None = None) -> dict:
    def _fake_proc(doc, st, bloque, prioridad, ocds_ctx):
        return fake[doc["sha256"]](doc)
    monkeypatch.setattr(D, "_procesar_doc", _fake_proc)
    state = state if state is not None else {"ocid": "1", "ocds": {"tender": {"description": "ADQUISICIÓN"}}}
    D.parse_documentos_lote(state, docs)
    return state


# ── 1225058: marca y precio de la OC no contaminan el ítem de las bases ────────────────
def test_marca_y_precio_de_la_oc_van_a_items_contratados_y_se_cruzan(monkeypatch):
    bases = _doc("b", "Bases Administrativas", "biddingDocuments", "sha_bases")
    oc = _doc("oc", "Archivos del contrato", "contractSigned", "sha_oc", seccion="contract")
    pag_bases = ["REQUERIMIENTO\nARROZ SUPERIOR 25,450 KG en sacos de 50 kg. Registro sanitario DIGESA."]
    pag_oc = ["ORDEN DE COMPRA N° 0001577\nARROZ SUPERIOR - SOMOS DEL NORTE | 25,450 KLG | 3.124126 | 79,509.00"]
    ext_bases = {"tipo_documento_detectado": "bases_administrativas", "contiene_requerimiento": True,
                 "items": [{"numero": "1", "descripcion_corta": "ARROZ SUPERIOR", "cantidad": 25450, "unidad": "KILOS",
                            "precio_unitario_referencial": None, "marca_o_modelo_exigido": None,
                            "texto_literal": "ARROZ SUPERIOR 25,450 KG en sacos de 50 kg. Registro sanitario DIGESA." * 2,
                            "evidencia": [{"pagina": 1, "cita": "ARROZ SUPERIOR 25,450 KG"}]}]}
    # La OC repite las EETT: el LLM marca contiene_requerimiento=True y pone la marca/precio
    # en los campos de "requerimiento" (así ocurrió en la corrida real, texto/702f2482).
    ext_oc = {"tipo_documento_detectado": "orden_de_compra", "contiene_requerimiento": True,
              "items": [{"numero": "1", "descripcion_corta": "ARROZ SUPERIOR - SOMOS DEL NORTE", "cantidad": 25450.0,
                         "unidad": "KLG", "precio_unitario_referencial": 3.124126, "cuantia_referencial_item": 79509.0,
                         "marca_o_modelo_exigido": "SOMOS DEL NORTE",
                         "texto_literal": "ARROZ SUPERIOR - SOMOS DEL NORTE 25,450 KLG 3.124126 79,509.00 " * 3,
                         "evidencia": [{"pagina": 1, "cita": "ARROZ SUPERIOR - SOMOS DEL NORTE | 25,450 KLG | 3.124126"}]}],
              "postores": [{"ruc": "20523905679", "razon_social": "COMERCIAL DELBUENO E.I.R.L.", "monto_oferta": 79509.0,
                            "es_ganador": True, "evidencia": [{"pagina": 1, "cita": "ORDEN DE COMPRA N° 0001577"}]}]}
    fake = {"sha_bases": lambda d: _res(d, ext_bases, pag_bases), "sha_oc": lambda d: _res(d, ext_oc, pag_oc)}
    state = _lote(monkeypatch, [bases, oc], fake)
    raw = state["parser_raw_consolidated"]

    assert len(raw["items_consolidados"]) == 1
    it = raw["items_consolidados"][0]
    assert it["marca_o_modelo_exigido"] is None                    # las bases NO exigen marca
    assert it["precio_unitario_referencial"] is None
    assert it["precio_unitario_ofertado"] == pytest.approx(3.124126)
    assert it["origen_precio"] == "orden_de_compra" and it["marca_ofertada"] == "SOMOS DEL NORTE"
    assert it["precio_ofertado_documento_sha256"] == "sha_oc" and it["precio_ofertado_pagina"] == 1
    assert it["documento_sha256"] == "sha_bases"
    assert len(raw["items_contratados"]) == 1
    c = raw["items_contratados"][0]
    assert c["precio_unitario_contratado"] == pytest.approx(3.124126) and c["marca_ofertada"] == "SOMOS DEL NORTE"
    assert c["documento_sha256"] == "sha_oc" and c["pagina"] == 1 and c["origen"] == "orden_de_compra"
    assert c["cantidad"] == 25450.0 and c["unidad"] == "KLG"
    # la OC quedó registrada como recorte de etapa contratación (no como requerimiento)
    rec = [r for r in state["recortes"] if r["limite"] == "solo_documentos_con_requerimiento"][0]
    assert rec["omitido"][0]["etapa"] == "contratacion"
    assert raw["descartes_parser"][0]["campo"] == "marca_o_modelo_exigido" and raw["descartes_parser"][0]["valor"] == "SOMOS DEL NORTE"
    # el precio contratado también en items_otros_documentos (compat R3)
    assert raw["items_otros_documentos"][0]["precio_contratado"] == pytest.approx(3.124126)
    assert state["documentos_texto"]["sha_oc"]["etapa"] == "contratacion"


# ── 1225392: marca inferida de códigos de parte (LEICA) no está en el texto de las bases ──
def test_marca_exigida_solo_si_el_texto_la_exige():
    doc = _doc("b", "Bases Administrativas", "biddingDocuments", "s1")
    tx = _tx(["ESTACION TOTAL de 5 segundos. Bateria GEB361, cargador GKL341, modulo AutoHeight."])
    data = {"tipo_documento_detectado": "bases_administrativas", "contiene_requerimiento": True,
            "items": [{"descripcion_corta": "ESTACION TOTAL", "marca_o_modelo_exigido": "LEICA", "evidencia": []}]}
    out = D._post_procesar(data, tx, "s1", doc)
    it = out["items"][0]
    assert it["marca_o_modelo_exigido"] is None and it["marca_no_respaldada"] == "LEICA"
    assert out["_descartes_parser"][0]["campo"] == "marca_o_modelo_exigido"

    # caso positivo: el texto dice "Marca: … o similar"
    tx2 = _tx(["Especificaciones: Marca Caterpillar o similar, modelo 320D."])
    data2 = {"tipo_documento_detectado": "especificaciones_tecnicas", "contiene_requerimiento": True,
             "items": [{"descripcion_corta": "EXCAVADORA", "marca_o_modelo_exigido": "Caterpillar o similar", "evidencia": []}]}
    it2 = D._post_procesar(data2, tx2, "s2", None)["items"][0]
    assert it2["marca_o_modelo_exigido"] == "Caterpillar o similar" and it2["marca_exigida_pagina"] == 1


def test_documento_de_contratacion_remapea_precio_y_marca_a_ofertado():
    tx = _tx(["ORDEN DE COMPRA 0000423 | ESTACION TOTAL - LEICA | 2 | 45,000.00 | 90,000.00"])
    data = {"tipo_documento_detectado": "orden_de_compra", "contiene_requerimiento": True,
            "items": [{"descripcion_corta": "ESTACION TOTAL - LEICA", "cantidad": 2, "precio_unitario_referencial": 45000.0,
                       "marca_o_modelo_exigido": "LEICA", "evidencia": [{"pagina": 1, "cita": "ESTACION TOTAL - LEICA"}]}]}
    it = D._post_procesar(data, tx, "s", None)["items"][0]
    assert it["precio_unitario_contratado"] == 45000.0 and it["precio_unitario_referencial"] is None
    assert it["marca_ofertada"] == "LEICA" and it["marca_o_modelo_exigido"] is None
    assert it["origen_precio"] == "orden_de_compra" and it["_etapa"] == "contratacion"
    # acta → ofertado
    it2 = D._post_procesar({"tipo_documento_detectado": "acta_buena_pro",
                            "items": [{"descripcion_corta": "X", "precio_unitario_referencial": 10.0, "evidencia": []}]},
                           tx, "s", None)["items"][0]
    assert it2["precio_unitario_ofertado"] == 10.0 and it2["origen_precio"] == "oferta_ganadora"
    # sin tipo detectado → decide el título/documentType del DocRef
    assert D._es_doc_contratacion(None, {"titulo": "Archivos del contrato", "tipo": "contractSigned"}) is True
    assert D._es_doc_contratacion("otro", {"titulo": "Bases Integradas", "tipo": "biddingDocuments"}) is False
    assert D._es_doc_contratacion("informe_sustento") is False and D._es_doc_contratacion("adenda") is True


# ── 1225392 / 1225379 / 1225030: postores con RUC repetido se fusionan ───────────────────
def test_postores_repetidos_se_fusionan_y_conservan_montos_del_acta(monkeypatch):
    prop = _doc("p", "Documentos de Presentación de Propuestas", "biddingDocuments", "sha_prop", formato="zip")
    acta = _doc("a", "Documentos de Otorgamiento de Buena Pro", "awardNotice", "sha_acta", seccion="award", formato="zip")
    pag_prop = ["1 ABEL CARPIO COBOS 10426100725 18/06/2026 14:49:46\n2 GREEN STONE INVERSIONES S.A.C 20600501381\n"
                "3 CONSTRUCTORA ANDINA PARA EL DESARROLLO COANDE S.A.C 20607620718"]
    pag_acta = ["FORMATO N° 11\n9.1 DETALLE DEL PRECIO DE LA OFERTA\n1 ABEL CARPIO COBOS 95,000.00\n"
                "2 GREEN STONE INVERSIONES S.A.C 97,960.00\n3 CONSTRUCTORA ANDINA PARA EL DESARROLLO COANDE S.A.C 99,940.00"]
    ext_prop = {"tipo_documento_detectado": "propuesta_economica",
                "postores": [{"ruc": "10426100725", "razon_social": "ABEL CARPIO COBOS", "evidencia": [{"pagina": 1, "cita": "ABEL CARPIO COBOS 10426100725"}]},
                             {"ruc": "20600501381", "razon_social": "GREEN STONE INVERSIONES S.A.C", "evidencia": []},
                             {"ruc": "20607620718", "razon_social": "CONSTRUCTORA ANDINA PARA EL DESARROLLO COANDE S.A.C", "evidencia": []}]}
    ext_acta = {"tipo_documento_detectado": "acta_buena_pro",
                "postores": [{"ruc": "10426100725", "razon_social": "CARPIO COBOS ABEL", "monto_oferta": 95000.0, "es_ganador": True,
                              "orden_prelacion": 1, "estado": "admitido", "puntaje": 100.0,
                              "evidencia": [{"pagina": 1, "cita": "1 ABEL CARPIO COBOS 95,000.00"}]},
                             {"ruc": "20600501381", "razon_social": "GREEN STONE INVERSIONES S.A.C", "monto_oferta": 97960.0,
                              "es_ganador": False, "orden_prelacion": 2, "estado": "admitido",
                              "evidencia": [{"pagina": 1, "cita": "2 GREEN STONE INVERSIONES S.A.C 97,960.00"}]},
                             {"ruc": "20607620718", "razon_social": "CONSTRUCTORA ANDINA PARA EL DESARROLLO COANDE S.A.C",
                              "monto_oferta": 99940.0, "es_ganador": False, "orden_prelacion": 3, "estado": "admitido", "evidencia": []}],
                "items": [{"descripcion_corta": "ADQUISICION DE EQUIPOS TOPOGRAFICOS", "cuantia_referencial_item": 95000.0}]}
    fake = {"sha_prop": lambda d: _res(d, ext_prop, pag_prop), "sha_acta": lambda d: _res(d, ext_acta, pag_acta)}
    state = {"ocid": "1225392", "ocds": {"tender": {"description": "EQUIPOS TOPOGRAFICOS", "value": {"amount": 95000.0},
                                                    "tenderers": [{"id": "PE-RUC-10426100725", "name": "CARPIO COBOS ABEL"}]},
                                         "awards": [{"suppliers": [{"id": "PE-RUC-10426100725"}]}]}}
    _lote(monkeypatch, [prop, acta], fake, state)   # el reporte de propuestas se procesa ANTES que el acta
    raw = state["parser_raw_consolidated"]

    posts = raw["postores"]
    assert posts is raw["postores_consolidados"] and len(posts) == 3
    by_ruc = {p["ruc"]: p for p in posts}
    assert by_ruc["10426100725"]["monto_oferta"] == 95000.0 and by_ruc["10426100725"]["es_ganador"] is True
    assert by_ruc["20600501381"]["monto_oferta"] == 97960.0 and by_ruc["20607620718"]["monto_oferta"] == 99940.0
    assert by_ruc["10426100725"]["puntaje"] == 100.0 and by_ruc["10426100725"]["estado"] == "admitido"
    assert by_ruc["10426100725"]["documento_sha256"] == "sha_acta" and by_ruc["10426100725"]["pagina"] == 1
    assert sorted(by_ruc["10426100725"]["fuentes"]) == ["sha_acta", "sha_prop"]
    assert by_ruc["10426100725"]["razon_social_ocds"] == "CARPIO COBOS ABEL"
    assert [o["monto"] for o in raw["ofertas"]] == [95000.0, 97960.0, 99940.0]
    assert [o["orden"] for o in raw["ofertas"]] == [1, 2, 3] and raw["ofertas"][0]["es_ganador"] is True
    assert raw["ofertas"][0]["fuente"] == "sha_acta" and raw["ofertas"][0]["ruc"] == "10426100725"
    # el acta no es fuente de requerimiento ni su cuantía pisa el referencial de las bases
    assert raw["items_consolidados"] == [] and raw.get("cuantia_total") is None
    assert all("_fuente_adjudicacion" not in p for p in posts)


def test_acta_procesada_antes_que_propuestas_no_pierde_montos():
    lista: list[dict] = []
    D._fusionar_postor(lista, {"ruc": "10426100725", "razon_social": "ABEL CARPIO COBOS", "monto_oferta": 95000.0,
                               "es_ganador": True}, True, "sha_acta")
    D._fusionar_postor(lista, {"ruc": "10426100725", "razon_social": "CARPIO COBOS ABEL", "monto_oferta": None,
                               "es_ganador": None}, False, "sha_prop")
    assert len(lista) == 1 and lista[0]["monto_oferta"] == 95000.0 and lista[0]["es_ganador"] is True
    assert lista[0]["documento_sha256"] == "sha_acta"
    # un dato del acta manda sobre uno previo de un documento no adjudicatorio
    lista2: list[dict] = []
    D._fusionar_postor(lista2, {"ruc": "10426100725", "razon_social": "ABEL CARPIO COBOS", "monto_oferta": 1.0}, False, "s1")
    D._fusionar_postor(lista2, {"ruc": "10426100725", "razon_social": "ABEL CARPIO COBOS", "monto_oferta": 95000.0}, True, "s2")
    assert lista2[0]["monto_oferta"] == 95000.0 and lista2[0]["documento_sha256"] == "s2"


# ── 1225090: RUC inválido (OCR) no crea un tercer postor; 1225058: corrección contra el OCDS ──
def test_ruc_invalido_se_fusiona_por_nombre_o_se_corrige_con_ocds():
    assert D.ruc_valido("20501887286") and not D.ruc_valido("20501867286")
    assert D.ruc_valido("20608886622") and not D.ruc_valido("20608885622")
    assert not D.ruc_valido("2050188728") and not D.ruc_valido("PE-RUC-x")
    lista: list[dict] = []
    D._fusionar_postor(lista, {"ruc": "20501887286", "razon_social": "DIAGNOSTICA PERUANA S.A.C.", "monto_oferta": 649800.0,
                               "es_ganador": True}, True, "sha_acta")
    D._fusionar_postor(lista, {"ruc": "20501867286", "razon_social": "DIAGNOSTICA PERUANA B.A.C.", "es_ganador": True},
                       False, "sha_fianza")
    assert len(lista) == 1 and lista[0]["ruc"] == "20501887286" and lista[0]["ruc_ocr"] == "20501867286"
    assert lista[0]["monto_oferta"] == 649800.0
    # Dialca: 20608885622 (OCR) → 20608886622 (tender.tenderers)
    conocidos = {"20608886622": "INVERSIONES DIALCA S.A.C."}
    lista2: list[dict] = []
    e = D._fusionar_postor(lista2, {"ruc": "20608885622", "razon_social": "INVERSIONES DIALCA SAC", "monto_oferta": 87000.0},
                           True, "sha_acta", conocidos)
    assert e["ruc"] == "20608886622" and e["ruc_ocr"] == "20608885622" and e["razon_social_ocds"] == "INVERSIONES DIALCA S.A.C."
    # RUC inválido sin corrección posible ni nombre parecido → postor nuevo con ruc None (no inventa RUC)
    e2 = D._fusionar_postor(lista2, {"ruc": "20999999999", "razon_social": "OTRA EMPRESA S.A.C."}, True, "sha_acta", conocidos)
    assert e2["ruc"] is None and e2["ruc_ocr"] == "20999999999" and len(lista2) == 2


def test_rucs_y_ganadores_desde_ocds():
    state = {"ocds": {"tender": {"tenderers": [{"id": "PE-RUC-20600501381", "name": "GREEN STONE"}]},
                      "parties": [{"identifier": {"scheme": "PE-RUC", "id": "10426100725"}, "roles": ["tenderer", "supplier"], "name": "CARPIO"}],
                      "awards": [{"suppliers": [{"id": "PE-RUC-10426100725"}]}]}}
    assert D._rucs_ocds(state) == {"20600501381": "GREEN STONE", "10426100725": "CARPIO"}
    assert D._ganadores_ocds(state) == {"10426100725"}


# ── 1225379: lista de invitados (Comparación de Precios) y ganador no invitado ───────────
def test_lista_invitados_y_ganador_no_invitado(monkeypatch):
    bases = _doc("b", "Bases Administrativas", "biddingDocuments", "sha_b")
    acta = _doc("a", "Documentos de Otorgamiento de Buena Pro", "awardNotice", "sha_a", seccion="award")
    ext_b = {"tipo_documento_detectado": "bases_administrativas", "contiene_requerimiento": True,
             "items": [{"descripcion_corta": "POSTES DE CONCRETO", "texto_literal": "Poste de concreto armado centrifugado " * 3, "evidencia": []}],
             "invitados": [{"ruc": "10445949499", "razon_social": "MOISES ROMERO VILLAFUERTE", "evidencia": [{"pagina": 1, "cita": "MOISES ROMERO VILLAFUERTE"}]},
                           {"ruc": "10615661959", "razon_social": "DALIA SEGURA ORTEGA", "evidencia": []},
                           {"ruc": "10430797366", "razon_social": "JHON ROGER FLORES TTITO", "evidencia": []}]}
    ext_a = {"tipo_documento_detectado": "acta_buena_pro",
             "postores": [{"ruc": "20601234567", "razon_social": "CONSTRUCTORA Y CONSULTORIA PUSAY SAC", "monto_oferta": 65000.0, "es_ganador": True, "evidencia": []},
                          {"ruc": "10615661959", "razon_social": "SEGURA ORTEGA DALIA", "monto_oferta": 68500.0, "es_ganador": False, "evidencia": []},
                          {"ruc": "10430797366", "razon_social": "FLORES TTITO JHON ROGER", "monto_oferta": 72650.0, "es_ganador": False, "evidencia": []}]}
    fake = {"sha_b": lambda d: _res(d, ext_b, ["FORMATO DE INVITACION\nMOISES ROMERO VILLAFUERTE"]),
            "sha_a": lambda d: _res(d, ext_a, ["acta"])}
    state = _lote(monkeypatch, [bases, acta], fake)
    raw = state["parser_raw_consolidated"]
    assert [i["ruc"] for i in raw["lista_invitados"]] == ["10445949499", "10615661959", "10430797366"]
    assert raw["lista_invitados"][0]["estado"] == "invitado" and raw["lista_invitados"][0]["pagina"] == 1
    by = {p["razon_social"]: p for p in raw["postores"]}
    pusay = by["CONSTRUCTORA Y CONSULTORIA PUSAY SAC"]
    assert pusay["invitado"] is False and pusay["estado"] == "no_invitado" and pusay["es_ganador"] is True
    assert by["SEGURA ORTEGA DALIA"]["invitado"] is True and by["SEGURA ORTEGA DALIA"]["estado"] == "admitido"
    assert len(raw["postores"]) == 3          # los invitados NO se cuentan como postores
    assert len(raw["ofertas"]) == 3 and raw["ofertas"][0]["postor"] == "CONSTRUCTORA Y CONSULTORIA PUSAY SAC"


def test_ganador_desde_ocds_si_ningun_documento_lo_marca(monkeypatch):
    acta = _doc("a", "Acta", "awardNotice", "sha_a", seccion="award")
    ext_a = {"tipo_documento_detectado": "acta_buena_pro",
             "postores": [{"ruc": "10426100725", "razon_social": "CARPIO COBOS ABEL", "monto_oferta": 95000.0, "evidencia": []},
                          {"ruc": "20600501381", "razon_social": "GREEN STONE", "monto_oferta": 97960.0, "evidencia": []}]}
    state = {"ocid": "1", "ocds": {"tender": {}, "awards": [{"suppliers": [{"id": "PE-RUC-10426100725"}]}]}}
    _lote(monkeypatch, [acta], {"sha_a": lambda d: _res(d, ext_a, ["x"])}, state)
    posts = {p["ruc"]: p for p in state["parser_raw_consolidated"]["postores"]}
    assert posts["10426100725"]["es_ganador"] is True and posts["10426100725"]["es_ganador_fuente"] == "ocds"
    assert posts["20600501381"].get("es_ganador") is None


# ── 1225266 / 1225030: PDF idéntico repetido → un solo OCR ─────────────────────────────
def test_zip_con_tres_copias_del_acta_produce_una_unidad_y_recortes():
    acta = _pdf(["ACTA DE BUENA PRO", "pagina 2"])
    entries = {"0Acta.pdf": acta, "2Cuadro de evaluacion economica.pdf": acta, "3Acta.pdf": acta,
               "1Reporte.pdf": _pdf(["reporte de evaluacion tecnica"])}
    unidades, recortes = D._expandir_contenedor(_zip(entries), "buena_pro.zip", PRIO_BIENES)
    assert len(unidades) == 2
    dup = [r for r in recortes if r["limite"] == "duplicado_sha256"]
    assert len(dup) == 2 and all("copia byte a byte" in r["omitido"] for r in dup)
    # ZIP anidado: el mismo acta dentro de otro ZIP tampoco se OCR-ea de nuevo
    entries2 = {"acta.pdf": acta, "anexos.zip": _zip({"copia/acta.pdf": acta})}
    unidades2, recortes2 = D._expandir_contenedor(_zip(entries2), "x.zip", PRIO_BIENES)
    assert len(unidades2) == 1 and [r["limite"] for r in recortes2] == ["duplicado_sha256"]


def test_misma_unidad_pdf_en_dos_documentos_del_lote_se_ocrea_una_vez(monkeypatch):
    pdf = _pdf(["acta de buena pro", "segunda pagina"])
    llamadas = []
    original = D._paginas_pymupdf

    def _contando(blob, offset):
        llamadas.append(offset)
        return original(blob, offset)
    monkeypatch.setattr(D, "_paginas_pymupdf", _contando)
    D._UNIT_CACHE.clear()
    tx1 = D._texto_de_unidades([D._unidad("zip1/acta.pdf", "pdf", data=pdf)])
    tx2 = D._texto_de_unidades([D._unidad("prefacio.txt", "paginas", paginas=[{"texto": "hola"}]),
                                D._unidad("zip2/cuadro.pdf", "pdf", data=pdf)])
    assert len(llamadas) == 1                                   # segundo documento: servido de la caché de unidades
    assert [p["n"] for p in tx2["paginas"]] == [1, 2, 3]        # renumerado con el offset del segundo doc
    assert tx2["paginas"][1]["archivo"] == "zip2/cuadro.pdf" and "acta de buena pro" in tx2["paginas"][1]["texto"]
    assert tx1["paginas"][0]["texto"] == tx2["paginas"][1]["texto"]


def test_mismo_sha_en_dos_docs_del_lote_usa_cache_en_memoria(monkeypatch):
    import base64
    pdf = _pdf(["bases integradas del proceso", "requerimiento"])
    url1, url2 = "https://x/a?fileCode=1", "https://x/b?fileCode=2"
    state = {"ocid": "1", "ocds": {"tender": {}}, "docs_b64": {url1: base64.b64encode(pdf).decode(), url2: base64.b64encode(pdf).decode()}}
    docs = [{"id": "d1", "url": url1, "gs": None, "tipo": "biddingDocuments", "titulo": "Bases", "seccion": "tender", "formato": "pdf", "sha256": None},
            {"id": "d2", "url": url2, "gs": None, "tipo": "biddingDocuments", "titulo": "Bases Integradas", "seccion": "tender", "formato": "pdf", "sha256": None}]
    ocr = []
    orig = D._texto_de_unidades
    monkeypatch.setattr(D, "_texto_de_unidades", lambda u: (ocr.append(1), orig(u))[1])
    monkeypatch.setattr(D, "_extraer_documento", lambda tx, label, bloque, ctx, hint: {"items": [], "resumen": "r", "_usos": [], "_recortes": []})
    monkeypatch.setattr(D, "_sha_por_url_get", lambda ocid, url: None)
    monkeypatch.setattr(D, "_sha_por_url_put", lambda *a, **k: None)
    monkeypatch.setattr(D, "_pg", lambda: (_ for _ in ()).throw(RuntimeError("sin BD")))
    D._TX_MEM.clear()
    monkeypatch.setattr(D, "PARSE_LOTE_WORKERS", 1)
    res = D.parse_documentos_lote(state, docs)
    assert res["n_ok"] == 2 and len(ocr) == 1                     # segundo doc: texto desde la caché de memoria
    assert res["n_cache_texto"] == 1
    shas = {d["sha256"] for d in res["documentos"]}
    assert len(shas) == 1


def test_docai_memo_por_sha(monkeypatch):
    from tools import docai
    monkeypatch.setattr(docai, "_PROCESSOR_ID", "proc")
    llamadas = []

    def _fake(pdf_bytes, mime_type="application/pdf", page_offset=0):
        llamadas.append(page_offset)
        return {"paginas": [{"n": 1, "texto": "hola", "chars": 4}, {"n": 2, "texto": "mundo", "chars": 5}],
                "texto": "x", "n_paginas": 2, "motor": "docai", "truncado": False, "recortes": []}
    monkeypatch.setattr(docai, "_extract_docai_uncached", _fake)
    docai._MEMO.clear()
    a = docai.extract_docai(b"%PDF-1.4 fake", page_offset=0)
    b = docai.extract_docai(b"%PDF-1.4 fake", page_offset=10)
    c = docai.extract_docai(b"%PDF-1.4 otro", page_offset=0)
    assert len(llamadas) == 2 and a["memo"] is False and b["memo"] is True and c["memo"] is False
    assert [p["n"] for p in b["paginas"]] == [11, 12] and "⟦p.12⟧" in b["texto"]


# ── 1225256 / 1225266: DOCX en orden del cuerpo, imágenes por r:embed, páginas reales ────
def _png(color: tuple[int, int, int]) -> bytes:
    import fitz
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
    pix.clear_with(color[0])
    return pix.tobytes("png")


def test_docx_imagenes_en_orden_del_cuerpo_y_paginas_por_salto():
    from docx import Document
    from docx.shared import Inches
    from docx.enum.text import WD_BREAK
    d = Document()
    d.add_paragraph("SECCIÓN III REQUERIMIENTO")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Ítem"
    t.rows[0].cells[1].text = "Cantidad"
    d.add_picture(io.BytesIO(_png((10, 10, 10))), width=Inches(1))     # → word/media/image1.png (primera en el cuerpo)
    d.add_paragraph("texto tras la primera imagen").add_run().add_break(WD_BREAK.PAGE)
    d.add_paragraph("PÁGINA DOS: penalidades")
    d.add_picture(io.BytesIO(_png((200, 200, 200))), width=Inches(1))  # → word/media/image2.png
    buf = io.BytesIO()
    d.save(buf)
    blob = buf.getvalue()
    # Renombrar en el ZIP: image1 ↔ image2 (y sus rels) → el orden de namelist() ya no es el del cuerpo
    src = zipfile.ZipFile(io.BytesIO(blob))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for info in src.infolist():
            data = src.read(info)
            name = info.filename
            if name == "word/media/image1.png":
                name = "word/media/image2.png"
            elif name == "word/media/image2.png":
                name = "word/media/image1.png"
            elif name == "word/_rels/document.xml.rels":
                txt = data.decode("utf-8").replace("media/image1.png", "media/TMP.png").replace("media/image2.png", "media/image1.png").replace("media/TMP.png", "media/image2.png")
                data = txt.encode("utf-8")
            z.writestr(name, data)
    blob2 = out.getvalue()
    with zipfile.ZipFile(io.BytesIO(blob2)) as z:
        orden, rels = D._docx_media_en_orden(z)
        assert orden == ["word/media/image2.png", "word/media/image1.png"]   # orden del cuerpo, no del ZIP
        assert z.read("word/media/image2.png") == src.read("word/media/image1.png")
    unidades = D._docx_a_unidades(blob2, "bases.docx")
    txt = [u for u in unidades if u["kind"] == "paginas"][0]
    assert len(txt["paginas"]) == 2                                          # salto de página real → 2 páginas
    p1, p2 = txt["paginas"][0]["texto"], txt["paginas"][1]["texto"]
    assert p1.startswith("SECCIÓN III REQUERIMIENTO\nÍtem | Cantidad")         # tabla intercalada donde estaba
    assert "[imagen 1: word/media/image2.png" in p1 and "[imagen 2: word/media/image1.png" in p2
    assert "PÁGINA DOS" in p2
    imgs = [u for u in unidades if u["kind"] == "imagenes"][0]
    assert imgs["orden_imagenes"] == ["word/media/image2.png", "word/media/image1.png"]
    assert D._n_paginas_pdf(imgs["data"]) == 2


# ── 1225416 / 1225058: folio impreso vs índice real ──────────────────────────────────────
def test_folio_impreso_se_corrige_a_pagina_real_y_se_anota():
    pags = [f"texto de relleno {i}\nPágina {i - 1} de 69" for i in range(1, 26)]
    pags[23] = "Lugar de entrega: Calle Nueva Alta s/n, Huarcacanto.\nPágina 22 de 69"   # índice real 24, folio 22
    tx = _tx(pags)
    data = {"tipo_documento_detectado": "bases_administrativas",
            "items": [{"descripcion_corta": "ARROZ", "texto_literal_paginas": [22, 23],
                       "evidencia": [{"pagina": 22, "cita": "Lugar de entrega: Calle Nueva Alta s/n, Huarcacanto"}]}]}
    it = D._post_procesar(data, tx, "s", None)["items"][0]
    ev = it["evidencia"][0]
    assert ev["verificada"] is True and ev["pagina"] == 24 and ev["pagina_declarada"] == 22 and ev["folio"] == 22
    assert it["texto_literal_paginas"] == [24, 25] and it["texto_literal_paginas_declaradas"] == [22, 23]


def test_cita_por_bolsa_de_tokens_en_fila_reordenada():
    tx = _tx(["9.1 DETALLE DEL PRECIO DE LA OFERTA\n1 ABEL CARPIO COBOS\n2 GREEN STONE INVERSIONES S.A.C | 95,000.00\n"
              "3 CONSTRUCTORA ANDINA PARA EL DESARROLLO | 97,960.00"])
    data = {"postores": [{"razon_social": "GREEN STONE", "evidencia": [{"pagina": 1, "cita": "2 GREEN STONE INVERSIONES S.A.C | 97,960.00"},
                                                                       {"pagina": 1, "cita": "marca Caterpillar exigida en bases"}]}]}
    evs = D._post_procesar(data, tx, "s", None)["postores"][0]["evidencia"]
    assert evs[0]["verificada"] is True and evs[0]["verificacion"] == "tokens"
    assert evs[1]["verificada"] is False


# ── 1225058: bloque contrato (ampliación improcedente); 1225090: procedimiento_seleccion ─
def test_bloque_contrato_y_procedimiento_seleccion_consolidados(monkeypatch):
    bases = _doc("b", "Bases Administrativas", "biddingDocuments", "sha_b")
    pliego = _doc("p", "Pliego de absolución de consultas", "clarifications", "sha_p")
    acta = _doc("a", "Acta de buena pro", "awardNotice", "sha_a", seccion="award")
    res_amp = _doc("r", "Archivos de la ampliación del contrato", "contractAnnexe", "sha_r", seccion="contract")
    ext_b = {"tipo_documento_detectado": "bases_administrativas", "contiene_requerimiento": True, "cuantia_reservada": True,
             "items": [{"descripcion_corta": "REACTIVOS", "texto_literal": "kit de reactivos de inmunohematologia " * 3, "evidencia": []}],
             "procedimiento_seleccion": {"factores_evaluacion": [{"factor": "precio", "puntaje_max": 40, "pagina": 40},
                                                                 {"factor": "garantia comercial", "puntaje_max": 10, "pagina": 41}],
                                         "evidencia": []}}
    ext_p = {"tipo_documento_detectado": "absolucion_consultas",
             "procedimiento_seleccion": {"consultas_observaciones": [
                 {"n": 12, "postor": "DIAGNOSTICA PERUANA S.A.C.", "tema": "incluir mejora 2 titulacion automatica 1/2048",
                  "absuelta": "se_acoge_parcialmente", "cambio_en_bases": "Mejora 2: 3 puntos", "pagina": 12}],
                 "modificaciones_integracion": [{"texto_original": "capacitacion 10", "texto_integrado": "capacitacion 7",
                                                  "a_pedido_de": "DIAGNOSTICA PERUANA S.A.C.", "pagina": 14}], "evidencia": []}}
    ext_a = {"tipo_documento_detectado": "acta_buena_pro",
             "postores": [{"ruc": "20501887286", "razon_social": "DIAGNOSTICA PERUANA S.A.C.", "monto_oferta": 649800.0, "es_ganador": True, "puntaje": 96.78, "evidencia": []},
                          {"ruc": "20100070970", "razon_social": "SISTEMAS ANALITICOS SRL", "monto_oferta": 597510.0, "es_ganador": False, "puntaje": 90.0, "evidencia": []}],
             "procedimiento_seleccion": {"puntajes_por_postor": [
                 {"ruc": "20501887286", "razon_social": "DIAGNOSTICA PERUANA S.A.C.", "factor": "total", "puntaje": 96.78, "pagina": 2},
                 {"ruc": "20100070970", "razon_social": "SISTEMAS ANALITICOS SRL", "factor": "garantia comercial", "puntaje": 0, "pagina": 2}], "evidencia": []},
             "cuantia_total": 649800.0}
    ext_r = {"tipo_documento_detectado": "adenda",
             "ejecucion_contractual": {"ampliaciones_plazo": [{"n": 1, "dias_solicitados": 10, "solicitada_por": "COMERCIAL DELBUENO E.I.R.L.",
                                                                "fecha_solicitud": "2026-08-19", "resolucion": "Resolución OGA 506-2026",
                                                                "resultado": "improcedente", "motivo": "no acredita inicio/cese del hecho (art. 142 Reglamento)", "pagina": 1}],
                                       "penalidades_aplicadas": [], "adendas": [],
                                       "entregas": [{"n": 2, "fecha_prevista": "2026-08-21", "fecha_real": None, "cantidad": 11550, "pagina": 1}],
                                       "evidencia": [{"pagina": 1, "cita": "IMPROCEDENTE la ampliación de plazo"}]}}
    fake = {"sha_b": lambda d: _res(d, ext_b, ["bases"] * 45), "sha_p": lambda d: _res(d, ext_p, ["pliego"] * 20),
            "sha_a": lambda d: _res(d, ext_a, ["acta", "puntajes"]),
            "sha_r": lambda d: _res(d, ext_r, ["RESOLUCION OGA 506-2026: declarar IMPROCEDENTE la ampliación de plazo\nPágina 1 de 3"])}
    state = _lote(monkeypatch, [bases, pliego, acta, res_amp], fake)
    raw = state["parser_raw_consolidated"]

    ps = raw["procedimiento_seleccion"]
    assert [f["factor"] for f in ps["factores_evaluacion"]] == ["precio", "garantia comercial"]
    assert ps["factores_evaluacion"][0]["documento_sha256"] == "sha_b"
    assert ps["consultas_observaciones"][0]["absuelta"] == "se_acoge_parcialmente" and ps["consultas_observaciones"][0]["documento_sha256"] == "sha_p"
    assert ps["modificaciones_integracion"][0]["a_pedido_de"] == "DIAGNOSTICA PERUANA S.A.C."
    assert len(ps["puntajes_por_postor"]) == 2 and ps["puntajes_por_postor"][1]["puntaje"] == 0
    c = raw["contrato"]
    assert c["ampliaciones_plazo"][0]["resultado"] == "improcedente" and c["ampliaciones_plazo"][0]["documento_sha256"] == "sha_r"
    assert c["ampliaciones_plazo"][0]["folio"] == 1 and c["entregas"][0]["cantidad"] == 11550
    assert c["evidencia"][0]["verificada"] is True
    assert raw["cuantia_reservada"] is True
    # la cuantía del acta es el monto adjudicado, no el referencial
    assert raw.get("cuantia_total") is None and raw["monto_adjudicado_doc"] == 649800.0 and raw["monto_adjudicado_doc_sha256"] == "sha_a"
    assert [o["monto"] for o in raw["ofertas"]] == [597510.0, 649800.0] and raw["ofertas"][1]["es_ganador"] is True


# ── firmantes: misma persona con nombre abreviado; ítems con ruido OCR ─────────────────
def test_firmantes_misma_persona_y_cargos_distintos(monkeypatch):
    oc = _doc("oc", "Archivos del contrato", "contractSigned", "sha_oc", seccion="contract")
    res = _doc("r", "Resolución", "contractAnnexe", "sha_r", seccion="contract")
    ext_oc = {"tipo_documento_detectado": "orden_de_compra",
              "firmantes": [{"nombre_completo": "YHONY E. QUISPE CANAZA", "cargo": "Responsable de Almacén", "evidencia": []}]}
    ext_r = {"tipo_documento_detectado": "adenda",
             "firmantes": [{"nombre_completo": "Yhony Edwin Quispe Canaza", "cargo": "Jefe de la Oficina de Abastecimiento", "dni": "40000000", "evidencia": []},
                           {"nombre_completo": "Juan L. Callohuanca Burgos", "cargo": "Director OGA", "evidencia": []}]}
    state = _lote(monkeypatch, [oc, res], {"sha_oc": lambda d: _res(d, ext_oc, ["x"]), "sha_r": lambda d: _res(d, ext_r, ["y"])})
    fs = state["parser_raw_consolidated"]["firmantes"]
    assert len(fs) == 2
    f = fs[0]
    assert f["nombre_completo"] == "Yhony Edwin Quispe Canaza" and f["dni"] == "40000000"
    assert f["cargos"] == ["Responsable de Almacén", "Jefe de la Oficina de Abastecimiento"]
    assert f["documentos"] == ["Archivos del contrato", "Resolución"]
    assert not D._mismo_firmante({"nombre_completo": "PERCY BOLIVAR ESPINOZA"}, {"nombre_completo": "PERCY ZAPANA ILLACHURA"})


def test_item_key_tolera_ruido_ocr_y_lookalikes():
    assert D._desc_compacta("DRYWALL 0.90 mm") == "DRYWALL090MM"
    assert D._desc_compacta("DRYWALL0.90 MM Μ .") == "DRYWALL090MMM"
    it_a = {"descripcion_corta": "DRYWALL 0.90 mm", "cantidad": 120, "requerimiento_tecnico_detallado": "x" * 50}
    it_b = {"descripcion_corta": "DRYWALL0.90 MM Μ .", "cantidad": 120, "requerimiento_tecnico_detallado": "y" * 50}
    keys = {D._item_key(it_a): it_a}
    assert D._buscar_item_similar(keys, D._item_key(it_b)) is it_a
    it_c = {"descripcion_corta": "AMPLIFICADOR DE AUDIO DE 600 W", "cantidad": 120, "requerimiento_tecnico_detallado": "z" * 50}
    keys2 = {D._item_key({"descripcion_corta": "AMPLIFICADOR DE AUDIO", "cantidad": 120, "requerimiento_tecnico_detallado": "z" * 50}): {}}
    assert D._buscar_item_similar(keys2, D._item_key(it_c)) is None      # productos distintos no se funden


def test_cruce_items_contratados_por_raices_y_cantidad():
    cons = [{"descripcion_corta": "ESTACION TOTAL", "cantidad": 2}, {"descripcion_corta": "PRISMA PENTAGONAL CON PORTA PRISMA", "cantidad": 4},
            {"descripcion_corta": "TRIPODE DE ALUMINIO", "cantidad": 2}]
    contr = [{"descripcion": "PRISMA PENTAGONAL CON PORTA PRISMA - LEICA", "cantidad": 4, "precio_unitario_contratado": 990.0, "marca_ofertada": "LEICA", "origen": "orden_de_compra", "documento_sha256": "oc", "pagina": 1},
             {"descripcion": "ESTACION TOTAL - LEICA", "cantidad": 2, "precio_unitario_contratado": 45000.0, "marca_ofertada": "LEICA", "origen": "orden_de_compra", "documento_sha256": "oc", "pagina": 1},
             {"descripcion": "TRIPODES DE ALUMINIO - LEICA", "cantidad": 2, "precio_unitario_contratado": 520.0, "marca_ofertada": "LEICA", "origen": "orden_de_compra", "documento_sha256": "oc", "pagina": 1}]
    D._cruzar_items_contratados(cons, contr)
    assert [c["precio_unitario_ofertado"] for c in cons] == [45000.0, 990.0, 520.0]
    assert all(c["marca_ofertada"] == "LEICA" and c["origen_precio"] == "orden_de_compra" for c in cons)
    assert all(c.get("marca_o_modelo_exigido") is None for c in cons)
    # un contrato gana sobre una oferta para el mismo ítem
    cons2 = [{"descripcion_corta": "ARROZ SUPERIOR", "cantidad": 25450}]
    contr2 = [{"descripcion": "ARROZ SUPERIOR", "cantidad": 25450, "precio_unitario_contratado": 3.2, "origen": "oferta_ganadora", "documento_sha256": "a", "pagina": 2},
              {"descripcion": "ARROZ SUPERIOR - SOMOS DEL NORTE", "cantidad": 25450, "precio_unitario_contratado": 3.124126, "origen": "contrato", "documento_sha256": "c", "pagina": 1}]
    D._cruzar_items_contratados(cons2, contr2)
    assert cons2[0]["precio_unitario_ofertado"] == pytest.approx(3.124126) and cons2[0]["origen_precio"] == "contrato"


def test_sanitizador_no_rellena_requerimiento_desde_contratacion():
    """El merge de grupos del sanitizador (código, no LLM) nunca copia marca/precio referencial
    desde un ítem de contratación al de las bases."""
    base = {"descripcion_corta": "ARROZ SUPERIOR", "requerimiento_tecnico_detallado": "x" * 60, "marca_o_modelo_exigido": None,
            "precio_unitario_referencial": None, "_etapa": "requerimiento"}
    oc = {"descripcion_corta": "ARROZ SUPERIOR - SOMOS DEL NORTE", "requerimiento_tecnico_detallado": "y" * 100,
          "marca_o_modelo_exigido": "SOMOS DEL NORTE", "precio_unitario_referencial": 3.12, "marca_ofertada": "SOMOS DEL NORTE",
          "_etapa": "contratacion"}
    grp = [oc, base]
    grp.sort(key=lambda x: (x.get("_etapa") == "contratacion", -len(str(x.get("requerimiento_tecnico_detallado") or ""))))
    assert grp[0] is base
    out = dict(grp[0])
    for other in grp[1:]:
        for k, v in other.items():
            if other.get("_etapa") == "contratacion" and k in D._CAMPOS_SOLO_REQUERIMIENTO:
                continue
            if out.get(k) in (None, "", [], {}, 0) and v not in (None, "", [], {}, 0):
                out[k] = v
    assert out["marca_o_modelo_exigido"] is None and out["precio_unitario_referencial"] is None
    assert out["marca_ofertada"] == "SOMOS DEL NORTE"


# ── doc_select: la propuesta del postor no es "Bases" ─────────────────────────────────
def test_doc_select_propuesta_no_rankea_como_bases():
    assert categorias_de("Documentos de Presentación de Propuestas", "biddingDocuments") == ["propuesta"]
    assert categorias_de("Anexo 3", "biddingDocuments") == ["bases"]                 # sin título útil, manda el documentType
    assert categorias_de("Bases Integradas", "biddingDocuments") == ["bases_integradas", "bases"]
    r_prop = rank_documento("Documentos de Presentación de Propuestas", "biddingDocuments", PRIO_BIENES)
    r_acta = rank_documento("Documentos de Otorgamiento de Buena Pro", "awardNotice", PRIO_BIENES)
    r_bases = rank_documento("Bases Administrativas", "biddingDocuments", PRIO_BIENES)
    assert r_bases[0] < r_acta[0] < r_prop[0] and r_prop[1] == "propuesta"
    ocds = {"tender": {"documents": [
        {"id": "d1", "title": "Documentos de Presentación de Propuestas", "documentType": "biddingDocuments", "url": "https://x/1"},
        {"id": "d2", "title": "Bases Administrativas", "documentType": "biddingDocuments", "url": "https://x/2"}]},
        "awards": [{"documents": [{"id": "a1", "title": "Documentos de Otorgamiento de Buena Pro", "documentType": "awardNotice", "url": "https://x/3"}]}]}
    elegidos, _ = seleccionar_documentos("1", ocds, {}, PRIO_BIENES, 12, gcs_rows=[])
    assert [d["id"] for d in elegidos] == ["d2", "a1", "d1"]


# ═══════════════════════════════════════════════════════════════════════════════════════
# LIVE (RUN_LIVE=1 · ADC · GOOGLE_GENAI_USE_VERTEXAI=1 · DOCAI_PROCESSOR_ID · PGHOST/PGPASSWORD):
# parse_documentos_lote sobre expedientes reales desde GCS (texto OCR cacheado en BD; la
# extracción se recalcula por el cambio de PARSER_SCHEMA_VERSION).
#   1225392: 3 ofertas 95 000 / 97 960 / 99 940 (Formato 11 p.4), ganador al 100 % de la cuantía.
#   1225058: bases sin marca; OC con marca "SOMOS DEL NORTE" y P.U. 3.124126.
# ═══════════════════════════════════════════════════════════════════════════════════════
def _live_state(ocid: str) -> tuple[dict, list[dict]]:
    import json
    from tools._core import _pg
    from tools.doc_select import seleccionar_documentos, PRIORIDAD_DEFAULT
    conn = _pg()
    try:
        cur = conn.cursor()
        cur.execute("SELECT ocds_payload FROM convocatorias WHERE ocid = %s", (ocid,))
        row = cur.fetchone()
    finally:
        conn.close()
    assert row, f"convocatoria {ocid} no está en la BD"
    ocds = row[0] if isinstance(row[0], dict) else json.loads(row[0])
    elegidos, omitidos = seleccionar_documentos(ocid, ocds, {}, PRIORIDAD_DEFAULT, 12)
    assert elegidos, "sin documentos en documentos_gcs"
    return {"ocid": ocid, "ocds": ocds}, elegidos


def _live_dump(nombre: str, obj) -> None:
    import json
    out = os.getenv("LIVE_DUMP_DIR")
    if out:
        os.makedirs(out, exist_ok=True)
        with open(os.path.join(out, nombre), "w", encoding="utf-8") as f:
            json.dump(obj, f, ensure_ascii=False, indent=1, default=str)


@pytest.mark.live
def test_live_1225392_postores_con_los_tres_montos():
    state, docs = _live_state("1225392")
    res = D.parse_documentos_lote(state, docs)
    raw = state["parser_raw_consolidated"]
    _live_dump("1225392_raw.json", {"resumen": res, "raw": raw, "recortes": state["recortes"]})
    print("\nLIVE[1225392]", {k: res[k] for k in ("n_docs", "n_ok", "n_cache_texto", "n_cache_extraccion", "n_items_consolidados",
                                                "n_items_contratados", "n_items_con_precio_ofertado", "n_postores", "n_ofertas",
                                                "n_descartes_parser", "segundos")})
    for p in raw["postores"]:
        print("  postor", p.get("ruc"), p.get("razon_social"), p.get("monto_oferta"), p.get("es_ganador"), p.get("estado"),
              p.get("documento_sha256", "")[:8], p.get("pagina"))
    for it in raw["items_consolidados"]:
        print("  item", it.get("numero"), it.get("descripcion_corta"), "| exigida:", it.get("marca_o_modelo_exigido"),
              "| ofertada:", it.get("marca_ofertada"), "| ref:", it.get("precio_unitario_referencial"),
              "| ofertado:", it.get("precio_unitario_ofertado"), it.get("origen_precio"))
    assert res["n_ok"] == res["n_docs"] >= 4
    montos = sorted(o["monto"] for o in raw["ofertas"])
    assert montos == [95000.0, 97960.0, 99940.0], montos
    ganador = [p for p in raw["postores"] if p.get("es_ganador")]
    assert len(ganador) == 1 and ganador[0]["ruc"] == "10426100725" and ganador[0]["monto_oferta"] == 95000.0
    assert len(raw["postores"]) == 3 and all(D.ruc_valido(p["ruc"]) for p in raw["postores"])
    # las bases no exigen la marca LEICA (solo códigos de parte); la marca vive en la OC
    assert all(not it.get("marca_o_modelo_exigido") for it in raw["items_consolidados"])
    assert any((c.get("marca_ofertada") or "").upper().startswith("LEICA") for c in raw["items_contratados"])
    assert any(it.get("precio_unitario_ofertado") for it in raw["items_consolidados"])


@pytest.mark.live
def test_live_1225058_bases_sin_marca_oc_con_marca():
    state, docs = _live_state("1225058")
    res = D.parse_documentos_lote(state, docs)
    raw = state["parser_raw_consolidated"]
    _live_dump("1225058_raw.json", {"resumen": res, "raw": raw, "recortes": state["recortes"]})
    print("\nLIVE[1225058]", {k: res[k] for k in ("n_docs", "n_ok", "n_cache_texto", "n_cache_extraccion", "n_items_consolidados",
                                                "n_items_contratados", "n_items_con_precio_ofertado", "n_postores", "n_ofertas",
                                                "n_descartes_parser", "segundos", "contrato")})
    for it in raw["items_consolidados"]:
        print("  item", it.get("numero"), it.get("descripcion_corta"), "| exigida:", it.get("marca_o_modelo_exigido"),
              "| ofertada:", it.get("marca_ofertada"), "| ref:", it.get("precio_unitario_referencial"),
              "| ofertado:", it.get("precio_unitario_ofertado"), it.get("origen_precio"), "| sha", (it.get("documento_sha256") or "")[:8])
    print("  contratados", [(c["descripcion"], c["precio_unitario_contratado"], c["marca_ofertada"], c["origen"]) for c in raw["items_contratados"]])
    print("  contrato", raw["contrato"])
    print("  postores", len(raw["postores"]), "ofertas", [o["monto"] for o in raw["ofertas"]][:12])
    assert res["n_ok"] == res["n_docs"] >= 4
    assert raw["items_consolidados"], "sin ítems de las bases"
    assert all(not it.get("marca_o_modelo_exigido") for it in raw["items_consolidados"])
    it = raw["items_consolidados"][0]
    assert it["documento_sha256"].startswith("3b87a7bd")            # el ítem viene de las bases
    assert (it.get("marca_ofertada") or "").upper() == "SOMOS DEL NORTE"
    assert it.get("precio_unitario_ofertado") == pytest.approx(3.124126, abs=1e-3)
    assert it.get("origen_precio") in ("orden_de_compra", "contrato")
    assert it.get("precio_unitario_referencial") in (None, pytest.approx(3.97, abs=0.05))
    ganador = [p for p in raw["postores"] if p.get("es_ganador")]
    assert ganador and ganador[0]["ruc"] == "20523905679" and ganador[0]["monto_oferta"] == 79509.0
    assert len(raw["ofertas"]) >= 10                                  # 12 lances en el acta (11-12 legibles)
    assert all(D.ruc_valido(p["ruc"]) for p in raw["postores"] if p.get("ruc"))
