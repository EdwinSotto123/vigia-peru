"""Tests OFFLINE del parser en lote (tools/documentos.py · WS D): expansión de contenedores
sin topes, numeración global de páginas, verificación de evidencia, consolidación y
registro de recortes. Sin red, sin Gemini, sin Document AI, sin BD (todo inyectado)."""
from __future__ import annotations

import io
import os
import sys
import zipfile

import pytest

_AGENT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _AGENT not in sys.path:
    sys.path.insert(0, _AGENT)

os.environ.pop("DOCAI_PROCESSOR_ID", None)  # fallback PyMuPDF/nativo en estos tests

import tools.documentos as D  # noqa: E402


def _pdf(textos: list[str]) -> bytes:
    import fitz
    doc = fitz.open()
    for t in textos:
        page = doc.new_page()
        page.insert_text((72, 72), t, fontsize=11)
    b = doc.tobytes()
    doc.close()
    return b


def _zip(entries: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for k, v in entries.items():
            z.writestr(k, v)
    return buf.getvalue()


# ── contenedores ──────────────────────────────────────────────────────────────────────
def test_zip_sin_topes_ordenado_por_prioridad_y_recorte_7z():
    entries = {f"anexo_{i}.pdf": _pdf([f"anexo {i}"]) for i in range(5)}
    entries["Bases Integradas.pdf"] = _pdf(["bases integradas p1", "bases integradas p2"])
    entries["Acta de buena pro.pdf"] = _pdf(["acta"])
    entries["planos.7z"] = b"7z\xbc\xaf\x27\x1c" + b"\x00" * 20
    entries["nested/absolucion.zip"] = _zip({"absolucion.pdf": _pdf(["absolucion"])})
    unidades, recortes = D._expandir_contenedor(_zip(entries), "bases.zip", ("bases integradas", "acta", "absolucion"))
    nombres = [u["nombre"] for u in unidades]
    assert len(unidades) == 8                                   # 5 anexos + integradas + acta + absolución (sin tope de 3)
    assert nombres[0].endswith("Bases Integradas.pdf")
    assert nombres[1].endswith("Acta de buena pro.pdf")
    assert nombres[2].endswith("absolucion.pdf")                # zip anidado, por prioridad
    assert [r["limite"] for r in recortes] == ["formato_no_soportado"] and "planos.7z" in recortes[0]["omitido"]


def test_docx_y_xlsx_a_paginas_de_texto():
    from docx import Document
    d = Document()
    d.add_paragraph("TÉRMINOS DE REFERENCIA del servicio de vigilancia")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Personal clave"
    t.rows[0].cells[1].text = "1 supervisor"
    buf = io.BytesIO()
    d.save(buf)
    unidades, recortes = D._expandir_contenedor(buf.getvalue(), "tdr.docx")
    assert recortes == [] and unidades and unidades[0]["kind"] == "paginas"
    assert "Personal clave | 1 supervisor" in unidades[0]["paginas"][0]["texto"]

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Presupuesto"
    ws.append(["Partida", "Metrado", "Precio"])
    ws.append(["01.01 Excavación", 120, 35.5])
    xb = io.BytesIO()
    wb.save(xb)
    unidades, recortes = D._expandir_contenedor(xb.getvalue(), "presupuesto.xlsx")
    assert recortes == [] and unidades[0]["kind"] == "paginas"
    assert "[hoja: Presupuesto]" in unidades[0]["paginas"][0]["texto"]
    assert "01.01 Excavación | 120 | 35.5" in unidades[0]["paginas"][0]["texto"]


def test_pdf_cifrado_y_formato_desconocido_son_recortes():
    import fitz
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "secreto")
    enc = doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="clave", owner_pw="clave")
    doc.close()
    u, r = D._expandir_contenedor(enc, "cifrado.pdf")
    assert u == [] and r[0]["limite"] == "pdf_cifrado"
    u, r = D._expandir_contenedor(b"\x00\x01\x02\x03garbage", "raro.bin")
    assert u == [] and r[0]["limite"] == "formato_no_soportado"


# ── páginas globales con marcadores ───────────────────────────────────────────────────
def test_texto_de_unidades_numera_paginas_globalmente():
    unidades = [
        D._unidad("a.pdf", "pdf", data=_pdf(["hola pagina uno", "hola pagina dos"])),
        D._unidad("b.txt", "paginas", paginas=[{"texto": "texto nativo"}]),
    ]
    tx = D._texto_de_unidades(unidades)
    assert tx["n_paginas"] == 3 and [p["n"] for p in tx["paginas"]] == [1, 2, 3]
    assert tx["paginas"][2]["archivo"] == "b.txt"
    assert "⟦archivo: a.pdf⟧" in tx["texto"] and "⟦p.3⟧\ntexto nativo" in tx["texto"]
    assert "hola pagina dos" in tx["paginas"][1]["texto"]
    assert D._texto_rango(tx["paginas"], 2, 3).startswith("⟦archivo: a.pdf⟧\n⟦p.2⟧")


# ── evidencia verificable + alias ─────────────────────────────────────────────────────
def test_post_procesar_verifica_citas_y_estampa_sha():
    tx = {"paginas": [
        {"n": 1, "texto": "ÍTEM 1: SERVICIO DE VIGILANCIA\nPlazo de ejecución: 365 días calendario.", "chars": 60},
        {"n": 2, "texto": "Personal: 12 agentes con carné SUCAMEC vigente.", "chars": 40},
    ]}
    data = {"items": [{
        "descripcion_corta": "SERVICIO DE VIGILANCIA",
        "texto_literal": "x" * (D.TEXTO_LITERAL_MAX + 50),
        "evidencia": [
            {"pagina": 1, "cita": "Plazo de ejecución: 365 días calendario"},      # exacta
            {"pagina": 1, "cita": "12 agentes con carné SUCAMEC"},                  # está en p.2 → tolerancia ±1
            {"pagina": 2, "cita": "marca Caterpillar exigida"},                     # inventada
        ]}],
        "firmantes": [{"nombre_completo": "X", "evidencia": "no es lista"}],
    }
    out = D._post_procesar(data, tx, "sha_abc")
    it = out["items"][0]
    assert it["documento_sha256"] == "sha_abc"
    assert len(it["texto_literal"]) == D.TEXTO_LITERAL_MAX and it["texto_literal_truncado"] is True
    assert it["requerimiento_tecnico_detallado"] == it["texto_literal"]          # alias de compat
    evs = it["evidencia"]
    assert [e["verificada"] for e in evs] == [True, True, False]
    assert evs[1]["pagina"] == 2                                                  # corregida a la página real
    assert out["firmantes"][0]["evidencia"] == [] and out["firmantes"][0]["documento_sha256"] == "sha_abc"
    assert out["_evidencia_stats"] == {"total": 3, "verificadas": 2}


def test_merge_extraccion_une_rangos_sin_duplicar():
    a = {"items": [{"descripcion_corta": "A"}], "cuantia_total": None, "servicio": {"entregables": [{"nombre": "E1"}], "plazo_total_dias": None}}
    b = {"items": [{"descripcion_corta": "A"}, {"descripcion_corta": "B"}], "cuantia_total": 100.0,
         "servicio": {"entregables": [{"nombre": "E2"}], "plazo_total_dias": 30}}
    m = D._merge_extraccion(a, b)
    assert [i["descripcion_corta"] for i in m["items"]] == ["A", "B"]
    assert m["cuantia_total"] == 100.0
    assert [e["nombre"] for e in m["servicio"]["entregables"]] == ["E1", "E2"] and m["servicio"]["plazo_total_dias"] == 30


# ── perfil ─────────────────────────────────────────────────────────────────────────────
def test_perfil_params_defaults_dict_y_objeto():
    assert D._perfil_params({}) == (None, D.PRIORIDAD_DEFAULT, D.MAX_DOCS_DEFAULT)
    st = {"perfil": {"parser_bloque": "obra", "doc_prioridad": ("expediente tecnico",), "parse_max_docs": 5}}
    assert D._perfil_params(st) == ("obra", ("expediente tecnico",), 5)

    class P:  # imita el dataclass Profile de P
        parser_bloque = "servicio"
        doc_prioridad = ("tdr", "bases")
        parse_max_docs = 7
    assert D._perfil_params({"perfil": P()}) == ("servicio", ("tdr", "bases"), 7)
    assert D._perfil_params({"perfil": {"parser_bloque": "inexistente"}})[0] is None   # bloque desconocido → sin bloque
    assert D._perfil_params({"perfil": P()}, parser_bloque="sustento_directa")[0] == "sustento_directa"


# ── consolidación del lote (con _procesar_doc simulado) ───────────────────────────────
def _fake_resultado(doc, ext, n_pag=3, chars=1000, error=None, recortes=None):
    return {"doc": doc, "sha256": doc.get("sha256"), "recortes": recortes or [], "tiempos": {"total_s": 1.0},
            "cache": {"texto": False, "extraccion": False}, "error": error,
            "tx": {"n_paginas": n_pag, "chars": chars, "motor": "docai", "truncado": False, "paginas": []},
            "ext": ext}


def test_parse_documentos_lote_consolida_y_registra_recortes(monkeypatch):
    docs = [
        {"id": "d1", "url": "https://x/bases", "gs": None, "tipo": "biddingDocuments", "titulo": "Bases", "seccion": "tender", "formato": "pdf", "sha256": "s1"},
        {"id": "d2", "url": "https://x/acta", "gs": None, "tipo": "awardNotice", "titulo": "Acta", "seccion": "award", "formato": "pdf", "sha256": "s2"},
        {"id": "d3", "url": "https://x/roto", "gs": None, "tipo": "otro", "titulo": "Roto", "seccion": "tender", "formato": "pdf", "sha256": "s3"},
    ]
    ext_bases = {
        "tipo_documento_detectado": "bases_administrativas", "contiene_requerimiento": True,
        "items": [{"numero": "1", "descripcion_corta": "CEMENTO", "cantidad": 10, "texto_literal": "cemento portland tipo I " * 5,
                   "evidencia": [{"pagina": 3, "cita": "cemento", "verificada": True}], "documento_sha256": "s1"}],
        "comite_evaluacion": [{"nombre_completo": "INVENTADO EN BASES"}],      # gate: bases no tiene comité
        "fundamento_legal": ["Ley 32069"], "cuantia_total": 5000.0, "resumen": "bases de cemento",
        "servicio": None, "_evidencia_stats": {"total": 1, "verificadas": 1}, "_usos": [],
    }
    ext_acta = {
        "tipo_documento_detectado": "acta_buena_pro", "contiene_requerimiento": False,
        "items": [{"numero": "1", "descripcion_corta": "ADQUISICIÓN DE CEMENTO"}],  # cabecera sin requerimiento → otros_documentos
        "postores": [{"razon_social": "Proveedor SAC", "ruc": None, "es_ganador": True},
                     {"razon_social": "PROVEEDOR S.A.C.", "ruc": "20123456789"}],
        "firmantes": [{"nombre_completo": "ANA PEREZ", "cargo": "Presidente", "evidencia": [{"pagina": 1, "cita": "ANA PEREZ", "verificada": True}]}],
        "comite_evaluacion": [{"nombre_completo": "ANA PEREZ", "rol": "presidente"}],
        "motivos_adjudicacion": [{"ganador_razon_social": "Proveedor SAC"}],
        "lugar_fecha_acta": {"lugar": "Lima", "fecha": "2026-05-01"},
        "fundamento_legal": ["Ley 32069", "DS 009-2025-EF"], "_evidencia_stats": {"total": 1, "verificadas": 1}, "_usos": [],
    }
    fake = {"s1": ext_bases, "s2": ext_acta}

    def _fake_proc(doc, state, bloque, prioridad, ocds_ctx):
        if doc["sha256"] == "s3":
            return _fake_resultado(doc, {}, error="download_failed: 403",
                                   recortes=[{"donde": "descarga:Roto", "limite": "descarga_fallida", "omitido": "Roto (403)"}])
        return _fake_resultado(doc, fake[doc["sha256"]])

    monkeypatch.setattr(D, "_procesar_doc", _fake_proc)
    state = {"ocid": "1", "ocds": {"tender": {"description": "ADQUISICIÓN DE CEMENTO"}}, "perfil": {"parser_bloque": "servicio"}}
    res = D.parse_documentos_lote(state, docs)

    raw = state["parser_raw_consolidated"]
    assert res["n_docs"] == 3 and res["n_ok"] == 2 and res["n_error"] == 1
    assert [i["descripcion_corta"] for i in raw["items_consolidados"]] == ["CEMENTO"]
    assert len(raw["items_otros_documentos"]) == 1                      # no se pierde, se aparta
    assert len(raw["postores_consolidados"]) == 1 and raw["postores_consolidados"][0]["ruc"] == "20123456789"  # dedupe por nombre
    assert [f["nombre_completo"] for f in raw["firmantes_consolidados"]] == ["ANA PEREZ"] and raw["firmantes"] is raw["firmantes_consolidados"]
    assert raw["comite_evaluacion"] == [{"nombre_completo": "ANA PEREZ", "rol": "presidente"}]   # solo del acta
    assert raw["lugar_fecha_acta"]["lugar"] == "Lima" and raw["cuantia_total"] == 5000.0
    assert raw["fundamento_legal"] == ["Ley 32069", "DS 009-2025-EF"]
    assert raw["requerimiento_disponible"] is True and "bases de cemento" in raw["resumen_ejecutivo"]
    assert len(raw["documentos"]) == 3 and raw["documentos"][2]["error"].startswith("download_failed")
    limites = [r["limite"] for r in state["recortes"]]
    assert "descarga_fallida" in limites
    assert "solo_documentos_con_requerimiento" in limites              # ítems del acta apartados (antes: gate mudo)
    assert "solo_actas_cuadros_contratos" in limites                   # comité inventado en Bases descartado, registrado
    assert state["documentos_texto"]["s1"]["n_paginas"] == 3 and state["documentos_texto"]["s3"]["truncado"] is False
    assert state["document_analysis"]["_source"] == "parse_documentos_lote"
    assert state["_parsed_doc_cache"]["https://x/bases"]["n_items_consolidados"] == 1
    assert res["evidencia"] == {"total": 2, "verificadas": 2}


def test_parse_documentos_lote_sin_docs_registra_recorte():
    state = {}
    res = D.parse_documentos_lote(state, [])
    assert res["n_docs"] == 0 and state["recortes"][0]["limite"] == "sin_documentos"


# ── JSON truncado → re-pedir por rango de páginas ──────────────────────────────────────
def test_extraer_rango_reintenta_por_mitades_y_registra_recorte(monkeypatch):
    paginas = [{"n": n, "texto": f"pagina {n} item {n}", "chars": 20} for n in range(1, 9)]
    llamadas = []

    def _fake_llm(texto, label, bloque, ocds_ctx, rango, tipo_hint):
        llamadas.append(rango)
        ns = [int(x) for x in __import__("re").findall(r"⟦p\.(\d+)⟧", texto)]
        # rangos de > 2 páginas "se truncan"; los chicos salen enteros
        trunc = len(ns) > 2
        return {"items": [{"descripcion_corta": f"item {n}"} for n in ns[:2 if trunc else len(ns)]]}, trunc, \
            {"modelo": "x", "segundos": 0.1, "finish_reason": "MAX_TOKENS" if trunc else "STOP", "tokens_prompt": 1, "tokens_output": 1, "tokens_thoughts": 0}

    monkeypatch.setattr(D, "_llamar_extractor", _fake_llm)
    recortes, usos = [], []
    data = D._extraer_rango(paginas, 1, 8, "doc", None, {}, None, recortes, usos)
    # 1-8 trunca → 1-4 y 5-8 truncan → 1-2, 3-4, 5-6, 7-8 salen enteros (depth 2)
    assert llamadas[0] is None and len(llamadas) == 7
    assert sorted(i["descripcion_corta"] for i in data["items"]) == sorted(f"item {n}" for n in range(1, 9))
    assert data["_truncado"] is False and recortes == [] and len(usos) == 7

    # con depth agotado (páginas 1-16 → 1-8 → 1-4 aún trunca en depth 2) queda recorte y _truncado
    llamadas.clear(); recortes.clear(); usos.clear()
    paginas16 = [{"n": n, "texto": f"pagina {n}", "chars": 20} for n in range(1, 17)]
    data = D._extraer_rango(paginas16, 1, 16, "doc", None, {}, None, recortes, usos)
    assert data["_truncado"] is True
    assert recortes and all(r["limite"].startswith("max_output_tokens") for r in recortes)


def test_extraer_documento_parte_por_chars(monkeypatch):
    monkeypatch.setattr(D, "PARSE_MAX_CHARS_POR_LLAMADA", 50)
    paginas = [{"n": n, "texto": "x" * 30, "chars": 30} for n in range(1, 5)]   # 120 chars → 4 rangos de ≤50
    rangos = []

    def _fake_llm(texto, label, bloque, ocds_ctx, rango, tipo_hint):
        rangos.append(rango)
        return {"items": [{"descripcion_corta": f"r{rango}"}], "cuantia_total": 1.0}, False, \
            {"modelo": "x", "segundos": 0.1, "finish_reason": "STOP", "tokens_prompt": 1, "tokens_output": 1, "tokens_thoughts": 0}

    monkeypatch.setattr(D, "_llamar_extractor", _fake_llm)
    ext = D._extraer_documento({"paginas": paginas, "chars": 120}, "doc", None, {}, None)
    assert rangos == [(1, 1), (2, 2), (3, 3), (4, 4)]
    assert len(ext["items"]) == 4 and ext["_truncado"] is False and len(ext["_usos"]) == 4


# ── tool legacy parse_document_pdf: sigue funcionando sobre la expansión sin topes ─────
def test_parse_document_pdf_legacy_usa_expansion_sin_topes(monkeypatch):
    import base64
    entries = {f"doc{i}.pdf": _pdf([f"documento {i}"]) for i in range(5)}
    entries["cifrado.7z"] = b"7z\xbc\xaf\x27\x1c" + b"\x00" * 8
    url = "https://prod1.seace.gob.pe/x?fileCode=abc"
    state = {"docs_b64": {url: base64.b64encode(_zip(entries)).decode()}}
    vistos = []

    def _fake_single(blob, label):
        vistos.append(label)
        return {"tipo_documento_detectado": "bases_administrativas", "contiene_requerimiento": True,
                "items": [{"descripcion_corta": label, "texto_literal": "spec " * 20}], "_source": label}

    monkeypatch.setattr(D, "_parse_single_pdf_with_gemini", _fake_single)

    class Ctx:
        def __init__(self, st):
            self.state = st

    out = D.parse_document_pdf(url, Ctx(state))
    assert out["n_pdfs_procesados"] == 5 and out["n_items_consolidados"] == 5     # antes: tope de 3 PDF por ZIP
    assert state["recortes"][0]["limite"] == "formato_no_soportado"
    assert len(state["parser_raw_consolidated"]["items_consolidados"]) == 5
