"""
load_sancionados_osce.py · normaliza dataset/SANCIONADOS/reporte_sancionados.xlsx
(3 hojas: multa, definitivo, temporal) → tabla osce_sancionados (única).

Cuidados:
  · Todos los strings vienen con \xa0 (NBSP) al final — strip agresivo.
  · Hoja 'definitivo' trae filas SIN # ni RUC: son continuaciones de la fila previa
    aportando otra infracción. Las concatenamos al campo `infraccion` del padre.
  · RUC empieza con '10' → persona natural; con '20' → empresa.
  · Fechas DD/MM/YYYY (con o sin cero a la izquierda).

Uso:
  python backend/scripts/load_sancionados_osce.py \
      --xlsx dataset/SANCIONADOS/reporte_sancionados.xlsx \
      --dsn  "host=... port=5432 dbname=vigia user=postgres password=..." \
      --truncate
"""
from __future__ import annotations

import argparse
import sys
import unicodedata
from datetime import datetime
from decimal import Decimal, InvalidOperation

import openpyxl
import psycopg2
import psycopg2.extras


def clean(s) -> str:
    if s is None:
        return ""
    return str(s).replace("\xa0", " ").strip()


def normalize_name(s: str) -> str:
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return s.upper().strip()


def parse_date(s: str | None) -> datetime | None:
    s = clean(s)
    if not s:
        return None
    for fmt in ("%d/%m/%Y", "%-d/%-m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    parts = s.replace("-", "/").split("/")
    if len(parts) == 3:
        try:
            d, m, y = (int(p) for p in parts)
            return datetime(y, m, d).date()
        except ValueError:
            return None
    return None


def parse_money(s) -> Decimal | None:
    s = clean(s)
    if not s:
        return None
    try:
        return Decimal(s.replace(",", ""))
    except InvalidOperation:
        return None


def parse_hoja_multa(ws) -> list[dict]:
    """Cols: #, TIPO, Razon Social, RUC, Resolución, Fecha de Resolución,
              Monto de Multa, Infracción, Periodo de Suspensión, Desde, Hasta,
              Otra Infracción, Norma, Verificación de pago, Estado"""
    out, header = [], None
    for row in ws.iter_rows(values_only=True):
        if header is None:
            header = row
            continue
        if not row or all(c is None for c in row):
            continue
        if not clean(row[0]):
            continue
        razon = clean(row[2])
        ruc = clean(row[3])
        if not ruc:
            continue
        infr1 = clean(row[7])
        infr2 = clean(row[11])
        infraccion_full = "\n".join(filter(None, [infr1, infr2]))
        out.append({
            "tipo": "multa",
            "razon_social": razon,
            "razon_social_norm": normalize_name(razon),
            "ruc": ruc,
            "es_persona_natural": ruc.startswith("10"),
            "resolucion": clean(row[4]),
            "fecha_resolucion": parse_date(row[5]),
            "monto_multa_soles": parse_money(row[6]),
            "verificacion_pago": clean(row[13]) or None,
            "periodo": clean(row[8]) or None,
            "fecha_desde": parse_date(row[9]),
            "fecha_hasta": parse_date(row[10]),
            "infraccion": infraccion_full,
            "norma": clean(row[12]) or None,
            "estado": clean(row[14]) or None,
        })
    return out


def parse_hoja_inhab(ws, tipo: str) -> list[dict]:
    """Cols definitivo/temporal: #, TIPO, Razon Social, RUC, Resolución,
                                  Periodo, Desde, Hasta, Infracción,
                                  Otra Infracción, Norma, Estado"""
    out: list[dict] = []
    header = None
    last = None
    for row in ws.iter_rows(values_only=True):
        if header is None:
            header = row
            continue
        if not row or all(c is None for c in row):
            continue
        has_id = bool(clean(row[0]))
        if has_id:
            razon = clean(row[2])
            ruc = clean(row[3])
            if not ruc:
                continue
            infr1 = clean(row[8])
            infr2 = clean(row[9])
            infr = "\n".join(filter(None, [infr1, infr2]))
            entry = {
                "tipo": tipo,
                "razon_social": razon,
                "razon_social_norm": normalize_name(razon),
                "ruc": ruc,
                "es_persona_natural": ruc.startswith("10"),
                "resolucion": clean(row[4]),
                "fecha_resolucion": None,
                "monto_multa_soles": None,
                "verificacion_pago": None,
                "periodo": clean(row[5]) or None,
                "fecha_desde": parse_date(row[6]),
                "fecha_hasta": parse_date(row[7]),
                "infraccion": infr,
                "norma": clean(row[10]) or None,
                "estado": clean(row[11]) or None,
            }
            out.append(entry)
            last = entry
        else:
            extra = clean(row[8])
            if extra and last is not None:
                last["infraccion"] = (last["infraccion"] + "\n" + extra).strip()
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--dsn", required=False)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--truncate", action="store_true")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    rows: list[dict] = []
    rows += parse_hoja_multa(wb["multa"])
    rows += parse_hoja_inhab(wb["definitivo"], "definitivo")
    rows += parse_hoja_inhab(wb["temporal"], "temporal")
    wb.close()

    by_tipo: dict[str, int] = {}
    for r in rows:
        by_tipo[r["tipo"]] = by_tipo.get(r["tipo"], 0) + 1
    print(f"[parse] total={len(rows)}  por_tipo={by_tipo}", file=sys.stderr)

    if args.dry_run:
        for r in rows[:3] + rows[-2:]:
            print(r)
        return

    if not args.dsn:
        print("ERROR: --dsn requerido (o --dry-run)", file=sys.stderr)
        sys.exit(2)

    conn = psycopg2.connect(args.dsn)
    conn.autocommit = False
    with conn.cursor() as cur:
        if args.truncate:
            cur.execute("TRUNCATE osce_sancionados RESTART IDENTITY")
            print("[db] TRUNCATE ejecutado", file=sys.stderr)
        psycopg2.extras.execute_batch(cur, """
            INSERT INTO osce_sancionados
              (tipo, razon_social, razon_social_norm, ruc, es_persona_natural,
               resolucion, fecha_resolucion, monto_multa_soles, verificacion_pago,
               periodo, fecha_desde, fecha_hasta, infraccion, norma, estado)
            VALUES
              (%(tipo)s, %(razon_social)s, %(razon_social_norm)s, %(ruc)s, %(es_persona_natural)s,
               %(resolucion)s, %(fecha_resolucion)s, %(monto_multa_soles)s, %(verificacion_pago)s,
               %(periodo)s, %(fecha_desde)s, %(fecha_hasta)s, %(infraccion)s, %(norma)s, %(estado)s)
        """, rows, page_size=1000)
    conn.commit()
    print(f"[db] insertadas {len(rows)} filas en osce_sancionados", file=sys.stderr)
    conn.close()


if __name__ == "__main__":
    main()
